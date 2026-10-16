from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from itertools import islice

from openpyxl import load_workbook

from .xlsx_safety import MAX_IMPORT_ROWS, MAX_XLSX_ENTRIES, MAX_XLSX_EXPANDED_BYTES


MAX_SOURCE_BYTES = 16 * 1024 * 1024
MAX_COLUMNS = 128
MAX_CELLS = 2_000_000
MAX_SHEETS = 32


@dataclass(frozen=True)
class SourceCell:
    value: object
    formula: bool = False
    error: bool = False


def _xlsx(content: bytes, *, cached: bool):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_XLSX_ENTRIES or sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES:
            raise ValueError("XLSX 展开大小或条目数超过安全上限")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=cached)
    if len(workbook.sheetnames) > MAX_SHEETS:
        workbook.close()
        raise ValueError(f"工作表超过 {MAX_SHEETS} 个，请单独导出需要的表")
    return workbook


def source_sheets(content: bytes, name: str) -> list[str]:
    if not content or len(content) > MAX_SOURCE_BYTES:
        raise ValueError("成本文件为空或超过 16 MiB 上限")
    suffix = name.rsplit(".", 1)[-1].lower()
    if suffix == "xlsx":
        workbook = _xlsx(content, cached=False)
        try:
            return workbook.sheetnames
        finally:
            workbook.close()
    if suffix in {"csv", "tsv"}:
        if content.startswith(b"PK") or b"\x00" in content[:1024]:
            raise ValueError("文件内容与 CSV/TSV 格式不符，请重新导出文本表格")
        return ["文本表格"]
    raise ValueError("当前支持 XLSX、CSV、TSV；旧 XLS 请先另存为 XLSX")


def iter_source_rows(content: bytes, name: str, *, sheet: str, encoding: str = "utf-8-sig", delimiter: str | None = None, limit: int | None = None):
    sheets = source_sheets(content, name)
    if sheet not in sheets:
        raise ValueError("请选择文件中实际存在的工作表")
    suffix = name.rsplit(".", 1)[-1].lower()
    workbook = cached_workbook = None
    try:
        if suffix == "xlsx":
            workbook = _xlsx(content, cached=False)
            cached_workbook = _xlsx(content, cached=True)
            worksheet, cached_sheet = workbook[sheet], cached_workbook[sheet]
            worksheet.reset_dimensions()
            cached_sheet.reset_dimensions()
            rows = (
                tuple(SourceCell(cache.value if raw.data_type == "f" else raw.value, raw.data_type == "f", raw.data_type == "e" or (raw.data_type == "f" and cache.data_type == "e"))
                      for raw, cache in zip(raw_row, cached_row, strict=True))
                for raw_row, cached_row in zip(worksheet.iter_rows(), cached_sheet.iter_rows(), strict=True)
            )
        else:
            if encoding not in {"utf-8-sig", "gb18030"}:
                raise ValueError("请选择 UTF-8 或 GB18030 编码")
            delimiter = delimiter or ("\t" if suffix == "tsv" else ",")
            if delimiter not in {",", "\t", ";"}:
                raise ValueError("分隔符必须为逗号、制表符或分号")
            try:
                decoded = content.decode(encoding, errors="strict")
            except UnicodeError as exc:
                raise ValueError("文本编码读取失败，请选择正确编码后重新读取") from exc
            rows = (tuple(SourceCell(value) for value in row)
                    for row in csv.reader(io.StringIO(decoded, newline=""), delimiter=delimiter, strict=True))
        cells = 0
        selected = rows if limit is None else islice(rows, limit)
        for row_number, row in enumerate(selected, 1):
            if row_number > MAX_IMPORT_ROWS + 50:
                raise ValueError("源表行数超过安全上限")
            # Ignore unpopulated trailing cells, never ignore populated columns.
            end = len(row)
            while end and row[end - 1].value in (None, "") and not row[end - 1].formula:
                end -= 1
            row = row[:end]
            cells += len(row)
            if len(row) > MAX_COLUMNS or cells > MAX_CELLS:
                raise ValueError("表格列数或总单元格数超过安全上限")
            yield row_number, row
    except (csv.Error, zipfile.BadZipFile) as exc:
        raise ValueError("表格文件结构无效，请重新导出") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if cached_workbook is not None:
            cached_workbook.close()
