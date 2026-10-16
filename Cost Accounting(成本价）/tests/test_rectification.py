import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from cost_sync.catalog_import import preview_catalog_import
from cost_sync.cost_import import CostImportManager
from cost_sync.database import connect_database, copy_mapping_state, upsert_products
from cost_sync.excel import export_enabled_products
from cost_sync.local_api import ProductRepository
from cost_sync.models import Product
from cost_sync.publication import RecoveryRequiredError, publication_journal, publish_files, recover_publication, require_no_pending_publication
from cost_sync.shipping import calculate_shipping_fee, read_shipping_binding_page, save_shipping_product_binding, save_shipping_template, trial_shipping_fee
from cost_sync.update_history import read_update_history


class GenericCostRectificationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.database = self.root / "cost.sqlite3"
        self.excel = self.root / "cost.xlsx"
        with connect_database(self.database) as connection:
            upsert_products(connection, [Product("OLD", None, "保留商品", "5", 1, None, None)], synced_at="2026-10-16T08:00:00+08:00")
            connection.commit()
            export_enabled_products(connection, self.excel)
        self.manager = CostImportManager(self.database, self.excel)

    def tearDown(self):
        self.temporary.cleanup()

    def csv_options(self, text):
        source = self.root / "source.csv"
        source.write_text(text, encoding="utf-8-sig")
        layout = self.manager.inspect(str(source), {})
        return source, {"sheet": layout["sheet"], "header_row": layout["header_row"], "mapping": layout["proposed_mapping"],
                        "scope": layout["scope"], "source_label": "合成供应商", "same_catalog": True, "confirmed_basis": True, "mode": "merge"}

    def products(self):
        with connect_database(self.database) as connection:
            return {row["sku_id"]: row["cost_price"] for row in connection.execute("SELECT sku_id,cost_price FROM products")}

    def test_merge_preserves_absent_products_and_text_codes_with_commas(self):
        _, options = self.csv_options('商品编码,成本价\n"SKU,001",12.34\n0001,0\n')
        preview = self.manager.preview(options)
        self.assertTrue(preview["ready_to_import"])
        self.manager.apply(preview["preview_id"])
        self.assertEqual({"OLD": "5", "SKU,001": "12.34", "0001": "0"}, self.products())
        with self.assertRaises(ValueError):
            self.manager.apply(preview["preview_id"])

    def test_replace_requires_complete_snapshot_confirmation(self):
        _, options = self.csv_options("商品编码,成本价\nNEW,7\n")
        options["mode"] = "replace"
        before = self.database.read_bytes()
        with self.assertRaisesRegex(ValueError, "整个成本目录"):
            self.manager.preview(options)
        self.assertEqual(before, self.database.read_bytes())
        options["complete_snapshot"] = True
        preview = self.manager.preview(options)
        self.assertEqual(1, preview["removed"])
        self.manager.apply(preview["preview_id"])
        self.assertEqual({"NEW": "7"}, self.products())

    def test_selected_source_bytes_are_an_immutable_snapshot(self):
        source, options = self.csv_options("商品编码,成本价\nNEW,7\n")
        source.write_text("商品编码,成本价\nNEW,99\n", encoding="utf-8-sig")
        preview = self.manager.preview(options)
        self.manager.apply(preview["preview_id"])
        self.assertEqual("7", self.products()["NEW"])

    def test_database_change_invalidates_preview(self):
        _, options = self.csv_options("商品编码,成本价\nNEW,7\n")
        preview = self.manager.preview(options)
        with connect_database(self.database) as connection:
            connection.execute("UPDATE products SET cost_price='6' WHERE sku_id='OLD'")
        with self.assertRaisesRegex(ValueError, "数据已更新"):
            self.manager.apply(preview["preview_id"])
        self.assertEqual({"OLD": "6"}, self.products())

    def test_duplicate_and_ambiguous_money_do_not_write(self):
        _, options = self.csv_options('商品编码,成本价\nNEW,"1,234"\nNEW,8\n')
        preview = self.manager.preview(options)
        self.assertFalse(preview["ready_to_import"])
        self.assertEqual({"OLD": "5"}, self.products())

    def test_excel_error_type_is_rejected_but_identical_literal_text_is_valid(self):
        for error_cell in (True, False):
            with self.subTest(error_cell=error_cell):
                source = self.root / "typed.xlsx"
                workbook = Workbook()
                worksheet = workbook.active
                worksheet.append(["商品编码", "成本价"])
                worksheet.append(["#N/A", 7])
                worksheet["A2"].data_type = "e" if error_cell else "s"
                workbook.save(source)
                layout = self.manager.inspect(str(source), {})
                options = {"sheet": layout["sheet"], "header_row": layout["header_row"], "mapping": layout["proposed_mapping"],
                           "scope": layout["scope"], "source_label": "类型样本", "same_catalog": True, "confirmed_basis": True}
                preview = self.manager.preview(options)
                self.assertEqual(not error_cell, preview["ready_to_import"])

    def test_formulas_without_cached_results_are_not_zero_cost(self):
        source = self.root / "formula.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["商品编码", "成本价"])
        worksheet.append(["NEW", "=1+2"])
        workbook.save(source)
        layout = self.manager.inspect(str(source), {})
        preview = self.manager.preview({"sheet": layout["sheet"], "header_row": layout["header_row"], "mapping": layout["proposed_mapping"],
                                        "scope": layout["scope"], "source_label": "公式样本", "same_catalog": True, "confirmed_basis": True, "allow_cached_formulas": True})
        self.assertFalse(preview["ready_to_import"])

    def test_jushuitan_layout_can_move_without_timestamp_columns(self):
        ordinary = self.root / "ordinary.xlsx"
        combination = self.root / "combination.xlsx"
        for path, headers, row in (
            (ordinary, ["成本价", "商品状态", "商品名称", "商品编码"], [4, "启用", "普通", "ONE"]),
            (combination, ["组合商品编码", "组合商品名称", "组合成本价", "商品状态", "商品编码", "数量", "子商品成本价"], ["SET", "组合", 8, "启用", "ONE", 2, 4]),
        ):
            workbook = Workbook()
            workbook.active.title = "说明"
            workbook.active.append(["合成数据"])
            sheet = workbook.create_sheet("数据")
            sheet.append(["导出说明"])
            sheet.append(headers)
            sheet.append(row)
            workbook.save(path)
        preview = preview_catalog_import(ordinary, combination)
        self.assertEqual(2, preview.merged_products)
        self.assertTrue(preview.ready_to_import)

    def test_draft_and_saved_shipping_share_decimal_boundaries(self):
        template = {"name": "小数边界", "first_weight": "0.7", "first_fee": "3", "additional_weight": "0.1", "additional_fee": "1"}
        template_id = save_shipping_template(self.database, **template, changed_at="2026-10-16T08:30:00+08:00")
        draft = trial_shipping_fee(template, weight_kg="0.8")
        saved = calculate_shipping_fee(self.database, template_id=template_id, weight_kg="0.8")
        self.assertEqual(("4", 1), (draft.fee, draft.additional_steps))
        self.assertEqual(saved.fee, draft.fee)
        self.assertTrue(trial_shipping_fee(template, weight_kg=None).estimated_from_first_weight)

    def test_health_does_not_query_database(self):
        repository = ProductRepository(self.database)
        with patch.object(repository, "_connect", side_effect=AssertionError("health scanned database")):
            health = repository.health()
        self.assertTrue(health["ready"])
        self.assertIn("service_identity_v1", health["capabilities"])

    def test_gb18030_tsv_and_saved_profile_preserve_column_identity(self):
        source = self.root / "供应商.tsv"
        source.write_bytes("采购价\t商品编码\n7.5\t0002\n".encode("gb18030"))
        layout = self.manager.inspect(str(source), {"encoding": "gb18030"})
        options = {"sheet": layout["sheet"], "header_row": layout["header_row"], "mapping": layout["proposed_mapping"],
                   "scope": layout["scope"], "source_label": "供应商", "same_catalog": True, "confirmed_basis": True,
                   "encoding": "gb18030", "profile_name": "供应商格式"}
        preview = self.manager.preview(options)
        self.manager.apply(preview["preview_id"])
        self.assertEqual("7.5", self.products()["0002"])
        profiles = self.manager.inspect(None, {"encoding": "gb18030"})["profiles"]
        self.assertEqual("采购价", profiles[0]["mapping_headers"]["cost"])
        history = read_update_history(self.database)
        self.assertEqual("供应商.tsv", history[0]["source"]["file_name"])

    def test_legacy_pair_copy_preserves_declared_catalog_scope(self):
        _, options = self.csv_options("商品编码,成本价\nNEW,7\n")
        self.manager.apply(self.manager.preview(options)["preview_id"])
        with connect_database(self.root / "copy.sqlite3") as connection:
            copy_mapping_state(self.database, connection)
            self.assertEqual(options["scope"], connection.execute("SELECT value FROM sync_state WHERE key='catalog_scope'").fetchone()[0])

    def test_disabled_template_rejects_new_binding_and_binding_page_is_bounded(self):
        template_id = save_shipping_template(self.database, name="已停用", first_weight="1", first_fee="3", additional_weight="1", additional_fee="1", enabled=False, changed_at="2026-10-16T08:30:00+08:00")
        with self.assertRaisesRegex(ValueError, "停用"):
            save_shipping_product_binding(self.database, shop_name="合成店", douyin_product_id="P1", template_id=template_id, changed_at="2026-10-16T08:31:00+08:00")
        with connect_database(self.database) as connection:
            batch = connection.execute("INSERT INTO mapping_import_batches(source_file,source_sha256,total_rows,verified_rows,pending_rows,inserted_rows,existing_rows,status,imported_at) VALUES ('synthetic','hash',101,101,0,101,0,'applied','now')").lastrowid
            connection.executemany("INSERT INTO douyin_sku_mappings(douyin_sku_id,douyin_product_id,merchant_sku_code,jushuitan_sku_id,product_name,specification,status,import_batch_id,created_at,updated_at,shop_name) VALUES(?,?,?,?,?,'','verified',?,'now','now','合成店')", [(f'D{i}', f'P{i:03}', 'OLD', 'OLD', f'商品{i}', batch) for i in range(101)])
        first = read_shipping_binding_page(self.database, shop_name="合成店")
        second = read_shipping_binding_page(self.database, shop_name="合成店", page=2)
        self.assertEqual((100, 1, 101), (len(first["products"]), len(second["products"]), first["unbound_count"]))
        self.assertFalse({p["douyin_product_id"] for p in first["products"]} & {p["douyin_product_id"] for p in second["products"]})

    def test_missing_recovery_backup_disables_rollback_in_history(self):
        _, options = self.csv_options("商品编码,成本价\nNEW,7\n")
        result = self.manager.apply(self.manager.preview(options)["preview_id"])
        (self.root / result["backup_file"]).unlink()
        latest = read_update_history(self.database)[0]
        self.assertFalse(latest["can_rollback"])
        self.assertEqual("恢复备份不存在", latest["rollback_unavailable_reason"])


class PublicationRectificationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.database, self.excel = self.root / "cost.bin", self.root / "excel.bin"
        self.database.write_bytes(b"old-db")
        self.excel.write_bytes(b"old-excel")
        self.staged_db, self.staged_excel = self.root / "staged-db", self.root / "staged-excel"
        self.staged_db.write_bytes(b"new-db")
        self.staged_excel.write_bytes(b"new-excel")
        self.pairs = [(self.staged_db, self.database), (self.staged_excel, self.excel)]

    def tearDown(self):
        self.temporary.cleanup()

    def assert_old_pair(self):
        self.assertEqual(b"old-db", self.database.read_bytes())
        self.assertEqual(b"old-excel", self.excel.read_bytes())

    def test_interrupted_publication_is_recovered_on_next_start(self):
        real_replace = os.replace
        def interrupted(source, target):
            if Path(source) == self.staged_excel:
                raise KeyboardInterrupt("simulated process interruption")
            return real_replace(source, target)
        with patch("cost_sync.publication.os.replace", side_effect=interrupted):
            with self.assertRaises(KeyboardInterrupt):
                publish_files(self.pairs, database=self.database)
        with self.assertRaises(RecoveryRequiredError):
            require_no_pending_publication(self.database)
        recover_publication(self.database, self.excel)
        self.assert_old_pair()
        self.assertFalse(publication_journal(self.database).exists())

    def test_failed_compensation_preserves_backups_for_retry(self):
        real_replace = os.replace
        def interrupted(source, target):
            if Path(source) == self.staged_excel or (Path(target) == self.database and Path(source).suffix == ".restore"):
                raise PermissionError("simulated locked file")
            return real_replace(source, target)
        with patch("cost_sync.publication.os.replace", side_effect=interrupted):
            with self.assertRaises(RecoveryRequiredError):
                publish_files(self.pairs, database=self.database)
        self.assertEqual(2, len(list(self.root.glob("*.previous"))))
        recover_publication(self.database, self.excel)
        self.assert_old_pair()

    def test_commit_sync_error_and_partial_compensation_cannot_leave_committed_marker(self):
        real_replace, real_fsync = os.replace, os.fsync
        failed_sync = False
        journal = publication_journal(self.database)
        def fail_commit(fd):
            nonlocal failed_sync
            if journal.exists() and journal.read_bytes().endswith(b"committed\n") and not failed_sync:
                failed_sync = True
                raise OSError("simulated commit sync failure")
            return real_fsync(fd)
        def fail_database_restore(source, target):
            if Path(target) == self.database and Path(source).suffix == ".restore":
                raise PermissionError("simulated database lock")
            return real_replace(source, target)
        with patch("cost_sync.publication.os.fsync", side_effect=fail_commit), patch("cost_sync.publication.os.replace", side_effect=fail_database_restore):
            with self.assertRaises(RecoveryRequiredError):
                publish_files(self.pairs, database=self.database)
        self.assertFalse(journal.read_bytes().endswith(b"committed\n"))
        self.assertEqual(b"old-excel", self.excel.read_bytes())
        recover_publication(self.database, self.excel)
        self.assert_old_pair()

    def test_cleanup_error_does_not_turn_committed_publication_into_failure(self):
        real_unlink = Path.unlink
        def locked_backup(path, *args, **kwargs):
            if path.suffix == ".previous":
                raise PermissionError("simulated cleanup lock")
            return real_unlink(path, *args, **kwargs)
        with patch.object(Path, "unlink", locked_backup):
            publish_files(self.pairs, database=self.database)
        require_no_pending_publication(self.database)
        recover_publication(self.database, self.excel)
        self.assertEqual(b"new-db", self.database.read_bytes())
        self.assertEqual(b"new-excel", self.excel.read_bytes())

    def test_incomplete_journal_never_guesses_or_overwrites_data(self):
        publication_journal(self.database).write_bytes(b'{"version":1')
        with self.assertRaises(RecoveryRequiredError):
            recover_publication(self.database, self.excel)
        self.assert_old_pair()


if __name__ == "__main__":
    unittest.main()
