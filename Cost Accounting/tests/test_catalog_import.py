import tempfile
import unittest
import hashlib
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from cost_sync.catalog_import import apply_catalog_import, preview_catalog_import
from cost_sync.database import connect_database, get_sync_state, upsert_products
from cost_sync.mapping_import import preview_mapping_import
from cost_sync.models import Product
from cost_sync.update_history import read_update_history
from cost_sync.update_rollback import rollback_latest_update


ORDINARY_HEADERS = ("商品编码", "商品名称", "成本价", "商品状态", "创建时间", "修改时间")
COMBINATION_HEADERS = (
    "组合商品编码",
    "组合商品名称",
    "组合成本价",
    "商品状态",
    "创建时间",
    "修改时间",
    "商品编码",
    "数量",
    "子商品成本价",
)


class CatalogImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.ordinary = self.root / "ordinary.xlsx"
        self.combination = self.root / "combination.xlsx"
        self.database = self.root / "cost.sqlite3"
        self.output_excel = self.root / "enabled.xlsx"

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def write_book(self, path: Path, headers, rows, *, created: datetime | None = None) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        if created is not None:
            workbook.properties.created = created
        workbook.save(path)
        workbook.close()

    def write_valid_pair(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [
                ("SINGLE", "单品", 5, "启用", "2026/09/01", "2026/09/01"),
                ("PACK", "普通表组合规格", None, "启用", "2026/09/01", "2026/09/01"),
            ],
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [
                ("PACK", "两件组合", 12, "启用", "2026/09/01", "2026/09/01", "A", 1, 5),
                ("PACK", "两件组合", 12, "启用", "2026/09/01", "2026/09/01", "B", 1, 7),
            ],
        )

    def test_preview_folds_combination_rows_and_resolves_only_missing_ordinary_cost(self) -> None:
        self.write_valid_pair()

        preview = preview_catalog_import(self.ordinary, self.combination)

        self.assertTrue(preview.ready_to_import)
        self.assertEqual((2, 2), (preview.ordinary_rows, preview.combination_rows))
        self.assertEqual((2, 1, 2), (
            preview.ordinary_products,
            preview.combination_products,
            preview.merged_products,
        ))
        self.assertEqual(("PACK",), preview.resolved_overlaps)
        self.assertEqual((), preview.conflicting_overlaps)
        products = {product.sku_id: product for product in preview.products}
        self.assertEqual("12", products["PACK"].cost_price)
        self.assertEqual("两件组合", products["PACK"].name)
        self.assertEqual((), preview.combination_cost_mismatches)

    def test_preview_reports_combination_child_cost_mismatch(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [("SINGLE", "单品", 5, "启用", "2026/09/01", "2026/09/01")],
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [
                ("PACK", "组合", 12, "启用", "2026/09/01", "2026/09/01", "A", 1, 5),
                ("PACK", "组合", 12, "启用", "2026/09/01", "2026/09/01", "B", 1, 8),
            ],
        )

        preview = preview_catalog_import(self.ordinary, self.combination)

        self.assertTrue(preview.ready_to_import)
        self.assertEqual(("PACK",), preview.combination_cost_mismatches)

    def test_preview_rejects_catalog_files_from_different_snapshot_dates(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [("SINGLE", "单品", 5, "启用", "2026/09/01", "2026/09/01")],
            created=datetime(2026, 9, 1, 9, 0),
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [("PACK", "组合", 5, "启用", "2026/09/02", "2026/09/02", "SINGLE", 1, 5)],
            created=datetime(2026, 9, 2, 9, 0),
        )

        with self.assertRaisesRegex(ValueError, "同一个导出日期"):
            preview_catalog_import(self.ordinary, self.combination)

    def test_preview_ignores_product_update_dates_when_export_dates_match(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [("SINGLE", "单品", 5, "启用", "2026/08/01", "2026/09/07")],
            created=datetime(2026, 9, 5, 9, 0),
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [("PACK", "组合", 10, "启用", "2026/08/01", "2026/09/01", "SINGLE", 2, 5)],
            created=datetime(2026, 9, 5, 9, 1),
        )

        preview = preview_catalog_import(self.ordinary, self.combination)

        self.assertTrue(preview.ready_to_import)

    def test_excel_escape_in_catalog_code_is_decoded_before_exact_mapping(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [("SINGLE", "单品", 5, "启用", "2026/09/05", "2026/09/05")],
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [
                (
                    "PACK_x002B_TWO",
                    "组合商品",
                    7.4,
                    "启用",
                    "2026/09/05",
                    "2026/09/05",
                    "SINGLE",
                    1,
                    5,
                )
            ],
        )
        preview = preview_catalog_import(self.ordinary, self.combination)
        apply_catalog_import(
            self.database,
            self.output_excel,
            preview,
            confirmed_ordinary_sha256=preview.ordinary_sha256,
            confirmed_combination_sha256=preview.combination_sha256,
            imported_at="2026-09-05T12:00:00+08:00",
        )
        douyin = self.root / "douyin.xlsx"
        self.write_book(
            douyin,
            (
                "商品ID", "商品名称", "一级类目", "二级类目", "三级类目", "四级类目",
                "商家SKU编码", "规格ID（SKUID）", "商品规格",
            ),
            [
                (
                    "P-1", "组合商品", "食品饮料", "休闲食品", "组合食品", "",
                    "PACK+TWO", "D-1", "组合装",
                )
            ],
        )

        mapping = preview_mapping_import(self.database, douyin)

        self.assertEqual((1, 0), (mapping.verified_rows, mapping.pending_rows))
        self.assertEqual("PACK+TWO", mapping.candidates[0].jushuitan_sku_id)

    def test_catalog_codes_decode_other_observed_excel_escapes_once(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [
                ("RANGE_x00B1_5", "范围", 1, "启用", "2026/09/05", "2026/09/05"),
                ("SIZE_x00D7_2", "尺寸", 2, "启用", "2026/09/05", "2026/09/05"),
                ("LITERAL_x005F_x002B_", "字面转义", 3, "启用", "2026/09/05", "2026/09/05"),
            ],
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [("PACK", "组合", 4, "启用", None, None, "RANGE_x00B1_5", 1, 1)],
        )

        preview = preview_catalog_import(self.ordinary, self.combination)
        sku_ids = {product.sku_id for product in preview.products}

        self.assertIn("RANGE±5", sku_ids)
        self.assertIn("SIZE×2", sku_ids)
        self.assertIn("LITERAL_x002B_", sku_ids)

    def test_catalog_rejects_duplicate_codes_created_by_excel_escape_decoding(self) -> None:
        self.write_book(
            self.ordinary,
            ORDINARY_HEADERS,
            [
                ("PACK+TWO", "组合一", 1, "启用", "2026/09/05", "2026/09/05"),
                ("PACK_x002B_TWO", "组合二", 2, "启用", "2026/09/05", "2026/09/05"),
            ],
        )
        self.write_book(
            self.combination,
            COMBINATION_HEADERS,
            [("OTHER", "其它组合", 3, "启用", None, None, "PACK+TWO", 1, 1)],
        )

        with self.assertRaisesRegex(ValueError, "商品编码重复：PACK\\+TWO"):
            preview_catalog_import(self.ordinary, self.combination)

    def test_preview_rejects_workbook_over_row_limit(self) -> None:
        self.write_valid_pair()

        with patch("cost_sync.catalog_import.MAX_IMPORT_ROWS", 1):
            with self.assertRaisesRegex(ValueError, "数据行超过"):
                preview_catalog_import(self.ordinary, self.combination)

    def test_preview_blocks_overlap_when_both_tables_have_different_costs(self) -> None:
        self.write_valid_pair()
        workbook = load_workbook(self.ordinary)
        workbook.active["C3"] = 10
        workbook.save(self.ordinary)
        workbook.close()

        preview = preview_catalog_import(self.ordinary, self.combination)

        self.assertFalse(preview.ready_to_import)
        self.assertEqual(("PACK",), preview.conflicting_overlaps)

    def test_apply_publishes_complete_snapshot_preserves_mappings_and_creates_backup(self) -> None:
        self.write_valid_pair()
        with connect_database(self.database) as connection:
            upsert_products(
                connection,
                [Product("OLD", None, "旧商品", "1", 1, None, None)],
                synced_at="2026-08-31T00:00:00+08:00",
            )
            batch_id = connection.execute(
                """
                INSERT INTO mapping_import_batches (
                    source_file, source_sha256, total_rows, verified_rows,
                    pending_rows, inserted_rows, existing_rows, status, imported_at
                ) VALUES ('douyin.xlsx', 'hash', 1, 1, 0, 1, 0, 'applied',
                          '2026-08-31T09:00:00+08:00')
                """
            ).lastrowid
            connection.execute(
                """
                INSERT INTO douyin_sku_mappings (
                    douyin_sku_id, douyin_product_id, merchant_sku_code,
                    jushuitan_sku_id, product_name, specification,
                    category_level_1, category_level_2, category_level_3, category_level_4,
                    status, import_batch_id, created_at, updated_at
                ) VALUES ('D-1', 'P-1', 'SINGLE', 'SINGLE', '单品', '一件',
                          '食品饮料', '休闲食品', '坚果炒货', '',
                          'verified', ?, 'now', 'now')
                """,
                (batch_id,),
            )
            connection.commit()
        preview = preview_catalog_import(self.ordinary, self.combination)

        result = apply_catalog_import(
            self.database,
            self.output_excel,
            preview,
            confirmed_ordinary_sha256=preview.ordinary_sha256,
            confirmed_combination_sha256=preview.combination_sha256,
            imported_at="2026-09-01T12:00:00+08:00",
        )

        self.assertTrue(result.backup_path.is_file())
        self.assertEqual(2, result.total_count)
        with connect_database(self.database) as connection:
            self.assertEqual(
                [("PACK", "12"), ("SINGLE", "5")],
                [tuple(row) for row in connection.execute(
                    "SELECT sku_id, cost_price FROM products ORDER BY sku_id"
                )],
            )
            self.assertEqual(1, connection.execute(
                "SELECT COUNT(*) FROM douyin_sku_mappings"
            ).fetchone()[0])
            self.assertEqual(
                ("食品饮料", "休闲食品", "坚果炒货", ""),
                tuple(connection.execute(
                    "SELECT category_level_1, category_level_2, category_level_3, "
                    "category_level_4 FROM douyin_sku_mappings WHERE douyin_sku_id = 'D-1'"
                ).fetchone()),
            )
            self.assertEqual("excel_pair_import", get_sync_state(connection, "last_mode"))
            update = connection.execute(
                "SELECT update_key, backup_file, backup_sha256 FROM system_updates"
            ).fetchone()
            self.assertEqual("catalog:2026-09-01T12:00:00+08:00", update["update_key"])
            self.assertEqual(result.backup_path.name, update["backup_file"])
            self.assertEqual(
                hashlib.sha256(result.backup_path.read_bytes()).hexdigest(),
                update["backup_sha256"],
            )
        workbook = load_workbook(self.output_excel, read_only=True, data_only=True)
        try:
            self.assertEqual(3, workbook["启用商品成本价"].max_row)
        finally:
            workbook.close()

        rollback_latest_update(
            self.database,
            self.output_excel,
            "catalog:2026-09-01T12:00:00+08:00",
            rolled_back_at="2026-09-01T12:30:00+08:00",
        )
        with connect_database(self.database) as connection:
            self.assertEqual(
                [("OLD", "1")],
                [tuple(row) for row in connection.execute(
                    "SELECT sku_id, cost_price FROM products ORDER BY sku_id"
                )],
            )
        rolled_back = next(
            entry for entry in read_update_history(self.database)
            if entry["update_id"] == "catalog:2026-09-01T12:00:00+08:00"
        )
        self.assertEqual("rolled_back", rolled_back["status"])

    def test_apply_rechecks_both_confirmed_hashes_before_writing(self) -> None:
        self.write_valid_pair()
        preview = preview_catalog_import(self.ordinary, self.combination)

        with self.assertRaisesRegex(ValueError, "SHA-256"):
            apply_catalog_import(
                self.database,
                self.output_excel,
                preview,
                confirmed_ordinary_sha256="wrong",
                confirmed_combination_sha256=preview.combination_sha256,
                imported_at="2026-09-01T12:00:00+08:00",
            )

        self.assertFalse(self.database.exists())

    def test_header_only_pair_and_forged_empty_preview_cannot_publish(self) -> None:
        self.write_book(self.ordinary, ORDINARY_HEADERS, [])
        self.write_book(self.combination, COMBINATION_HEADERS, [])
        with self.assertRaisesRegex(ValueError, "没有商品数据"):
            preview_catalog_import(self.ordinary, self.combination)

        self.write_valid_pair()
        preview = preview_catalog_import(self.ordinary, self.combination)
        forged = replace(preview, ordinary_products=0, merged_products=0)
        with self.assertRaisesRegex(ValueError, "空快照"):
            apply_catalog_import(
                self.database,
                self.output_excel,
                forged,
                confirmed_ordinary_sha256=forged.ordinary_sha256,
                confirmed_combination_sha256=forged.combination_sha256,
                imported_at="2026-09-01T12:00:00+08:00",
            )
        self.assertFalse(self.database.exists())


if __name__ == "__main__":
    unittest.main()
