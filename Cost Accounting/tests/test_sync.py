import tempfile
import unittest
import hashlib
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from cost_sync.config import AppConfig, TEST_API_BASE_URL
from cost_sync.database import connect_database, count_products, get_sync_state
from cost_sync.models import Product
from cost_sync.sync import SyncService, _publish_pair


class MockClient:
    def __init__(self, pages=None, error=None) -> None:
        self.pages = pages or []
        self.error = error
        self.calls = []

    def iter_product_pages(self, **kwargs):
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        yield from self.pages


class SyncTests(unittest.TestCase):
    def make_config(self, root: Path) -> AppConfig:
        return AppConfig(
            environment="test",
            api_base_url=TEST_API_BASE_URL,
            app_key="key",
            app_secret="secret",
            access_token="token",
            sync_start_date=date(2026, 1, 1),
            window_days=7,
            output_root=root,
        )

    def test_full_sync_writes_database_and_enabled_excel(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            now = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            products = [
                Product("001", 1, "启用商品", "12.34", 1, None, None),
                Product("002", 2, "禁用商品", "8", -1, None, None),
            ]
            result = SyncService(
                config, MockClient([products]), now=lambda: now
            ).run_full()

            self.assertEqual(2, result.total_count)
            self.assertEqual(1, result.enabled_count)
            with connect_database(config.database_path) as connection:
                self.assertEqual(2, count_products(connection))
                self.assertEqual("full", get_sync_state(connection, "last_mode"))
            workbook = load_workbook(config.excel_path)
            try:
                worksheet = workbook["启用商品成本价"]
                self.assertEqual(2, worksheet.max_row)
                self.assertEqual("001", worksheet["A2"].value)
                self.assertEqual("@", worksheet["A2"].number_format)
            finally:
                workbook.close()

    def test_failed_sync_preserves_previous_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            now = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            initial = MockClient([[Product("001", 1, "商品", "1", 1, None, None)]])
            SyncService(config, initial, now=lambda: now).run_full()
            database_before = config.database_path.read_bytes()
            excel_before = config.excel_path.read_bytes()

            failing = MockClient(error=RuntimeError("模拟接口失败"))
            with self.assertRaises(RuntimeError):
                SyncService(config, failing, now=lambda: now).run_full()
            self.assertEqual(database_before, config.database_path.read_bytes())
            self.assertEqual(excel_before, config.excel_path.read_bytes())

    def test_incremental_sync_handles_rename_and_disable(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            first_time = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            initial = MockClient([[Product("OLD", 7, "商品", "5", 1, None, None)]])
            SyncService(config, initial, now=lambda: first_time).run_full()

            second_time = datetime(2026, 1, 1, 2, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            changed = MockClient([[Product("NEW", 7, "商品", "6", -1, None, None)]])
            result = SyncService(config, changed, now=lambda: second_time).run_incremental()

            self.assertEqual(1, result.fetched_count)
            self.assertEqual(1, result.total_count)
            self.assertEqual(0, result.enabled_count)
            self.assertEqual("modified", changed.calls[0]["date_field"])
            self.assertEqual(first_time - timedelta(minutes=5), changed.calls[0]["begin"])
            with connect_database(config.database_path) as connection:
                row = connection.execute(
                    "SELECT sku_id, cost_price, enabled FROM products"
                ).fetchone()
                self.assertEqual(("NEW", "6", -1), tuple(row))
                self.assertEqual("incremental", get_sync_state(connection, "last_mode"))
            workbook = load_workbook(config.excel_path)
            try:
                self.assertEqual(1, workbook["启用商品成本价"].max_row)
            finally:
                workbook.close()

    def test_full_sync_preserves_mapping_batches_and_verified_mappings(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            first_time = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            product = Product("ERP-001", 1, "商品", "5", 1, None, None)
            SyncService(config, MockClient([[product]]), now=lambda: first_time).run_full()
            with connect_database(config.database_path) as connection:
                batch_id = connection.execute(
                    """
                    INSERT INTO mapping_import_batches (
                        source_file, source_sha256, total_rows, verified_rows,
                        pending_rows, inserted_rows, existing_rows, status, imported_at
                    ) VALUES ('douyin.xlsx', 'hash-1', 1, 1, 0, 1, 0, 'applied',
                              '2026-01-01T01:30:00+08:00')
                    """
                ).lastrowid
                connection.execute(
                    """
                    INSERT INTO douyin_sku_mappings (
                        douyin_sku_id, douyin_product_id, merchant_sku_code,
                        jushuitan_sku_id, product_name, specification, status,
                        import_batch_id, created_at, updated_at
                    ) VALUES ('D-001', 'P-001', 'ERP-001', 'ERP-001',
                              '商品', '规格', 'verified', ?, 'now', 'now')
                    """,
                    (batch_id,),
                )
                connection.commit()

            second_time = datetime(2026, 1, 2, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            SyncService(config, MockClient([[product]]), now=lambda: second_time).run_full()

            with connect_database(config.database_path) as connection:
                mapping = connection.execute(
                    "SELECT douyin_sku_id, jushuitan_sku_id FROM douyin_sku_mappings"
                ).fetchone()
                self.assertEqual(("D-001", "ERP-001"), tuple(mapping))
                self.assertEqual(
                    "applied",
                    connection.execute(
                        "SELECT status FROM mapping_import_batches WHERE id = ?", (batch_id,)
                    ).fetchone()[0],
                )
                update = connection.execute(
                    "SELECT update_key, backup_file, backup_sha256 FROM system_updates"
                ).fetchone()
                backup = config.database_path.parent / update["backup_file"]
                self.assertEqual(
                    "catalog:2026-01-02T01:00:00+08:00", update["update_key"]
                )
                self.assertEqual(
                    hashlib.sha256(backup.read_bytes()).hexdigest(),
                    update["backup_sha256"],
                )

    def test_empty_full_sync_cannot_replace_existing_nonempty_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            now = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            initial = MockClient([[Product("001", 1, "商品", "1", 1, None, None)]])
            SyncService(config, initial, now=lambda: now).run_full()
            database_before = config.database_path.read_bytes()
            excel_before = config.excel_path.read_bytes()

            with self.assertRaisesRegex(ValueError, "拒绝覆盖"):
                SyncService(config, MockClient([[]]), now=lambda: now).run_full()
            self.assertEqual(database_before, config.database_path.read_bytes())
            self.assertEqual(excel_before, config.excel_path.read_bytes())

    def test_publish_and_rollback_sources_share_target_directory(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            config = self.make_config(Path(temp_dir))
            now = datetime(2026, 1, 1, 1, 0, tzinfo=ZoneInfo("Asia/Shanghai"))
            initial = MockClient([[Product("001", 1, "商品", "1", 1, None, None)]])
            SyncService(config, initial, now=lambda: now).run_full()

            real_replace = __import__("os").replace
            replace_calls = []
            failed_excel_publish = False

            def replace_with_excel_failure(source, target):
                nonlocal failed_excel_publish
                source_path = Path(source)
                target_path = Path(target)
                replace_calls.append((source_path, target_path))
                if target_path == config.excel_path and not failed_excel_publish:
                    failed_excel_publish = True
                    raise PermissionError("模拟 Excel 被占用")
                return real_replace(source, target)

            changed = MockClient([[Product("002", 2, "新商品", "2", 1, None, None)]])
            with patch(
                "cost_sync.sync.os.replace", side_effect=replace_with_excel_failure
            ):
                with self.assertRaises(PermissionError):
                    SyncService(config, changed, now=lambda: now).run_full()

            self.assertEqual(
                [config.database_path, config.excel_path, config.database_path],
                [target for _, target in replace_calls],
            )
            for source, target in replace_calls:
                self.assertEqual(target.parent, source.parent)

    def test_second_publish_failure_restores_both_previous_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source_database = root / "new.sqlite3"
            source_excel = root / "new.xlsx"
            target_database = root / "cost_accounting.sqlite3"
            target_excel = root / "enabled_product_costs.xlsx"
            source_database.write_bytes(b"new-db")
            source_excel.write_bytes(b"new-xlsx")
            target_database.write_bytes(b"old-db")
            target_excel.write_bytes(b"old-xlsx")

            real_replace = __import__("os").replace

            def replace_with_second_failure(source, target):
                if Path(source) == source_excel:
                    raise PermissionError("模拟 Excel 被占用")
                return real_replace(source, target)

            with patch("cost_sync.sync.os.replace", side_effect=replace_with_second_failure):
                with self.assertRaises(PermissionError):
                    _publish_pair(
                        (source_database, target_database),
                        (source_excel, target_excel),
                    )

            self.assertEqual(b"old-db", target_database.read_bytes())
            self.assertEqual(b"old-xlsx", target_excel.read_bytes())


if __name__ == "__main__":
    unittest.main()
