from __future__ import annotations

import hashlib
import os
import sqlite3
import threading
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from cost_sync.database import connect_database, upsert_products
from cost_sync.excel import export_enabled_products
from cost_sync.models import Product
from cost_sync.update_rollback import (
    backup_database,
    record_system_update,
    rollback_latest_update,
)


class UpdateRollbackTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary.name)
        self.database = self.directory / "cost_accounting.sqlite3"
        self.excel = self.directory / "enabled_product_costs.xlsx"
        self._write_product("1", "2026-09-05T08:00:00+08:00")
        with connect_database(self.database) as connection:
            export_enabled_products(connection, self.excel)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _write_product(self, cost: str, synced_at: str) -> None:
        with connect_database(self.database) as connection:
            connection.execute("DELETE FROM products")
            upsert_products(
                connection,
                (
                    Product(
                        sku_id="SKU-1",
                        autoid=None,
                        name="测试商品",
                        cost_price=cost,
                        enabled=1,
                        created=None,
                        modified=None,
                    ),
                ),
                synced_at=synced_at,
            )

    def _record_catalog_update(self, update_key: str, completed_at: str) -> Path:
        backup = self.directory / f"{update_key.replace(':', '-')}.sqlite3.bak"
        backup_hash = backup_database(self.database, backup)
        self._write_product("2", completed_at)
        with connect_database(self.database) as connection:
            record_system_update(
                connection,
                update_key=update_key,
                kind="catalog",
                reference_id=completed_at,
                completed_at=completed_at,
                backup_path=backup,
                backup_sha256=backup_hash,
                details={"enabled_count": 1},
            )
        return backup

    def _database_hash(self) -> str:
        return hashlib.sha256(self.database.read_bytes()).hexdigest()

    def test_latest_catalog_update_rolls_back_and_rebuilds_excel(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")

        result = rollback_latest_update(
            self.database,
            self.excel,
            key,
            rolled_back_at="2026-09-05T10:00:00+08:00",
        )

        with connect_database(self.database) as connection:
            product = connection.execute(
                "SELECT cost_price FROM products WHERE sku_id = 'SKU-1'"
            ).fetchone()
            audit = connection.execute(
                "SELECT status, recovery_backup_file, recovery_backup_sha256 "
                "FROM system_updates WHERE update_key = ?",
                (key,),
            ).fetchone()
            integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        workbook = load_workbook(self.excel, read_only=True, data_only=True)
        try:
            excel_cost = workbook["启用商品成本价"]["C2"].value
        finally:
            workbook.close()

        self.assertEqual("1", product["cost_price"])
        self.assertEqual("rolled_back", audit["status"])
        self.assertEqual("ok", integrity)
        self.assertEqual(1, excel_cost)
        self.assertTrue(result.rebuilt_excel)
        self.assertTrue(result.recovery_backup_path.is_file())
        self.assertEqual(result.recovery_backup_path.name, audit["recovery_backup_file"])
        self.assertEqual(result.recovery_backup_sha256, audit["recovery_backup_sha256"])

    def test_non_latest_update_is_rejected(self) -> None:
        first = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(first, "2026-09-05T09:00:00+08:00")
        second_backup = self.directory / "second.sqlite3.bak"
        second_hash = backup_database(self.database, second_backup)
        with connect_database(self.database) as connection:
            record_system_update(
                connection,
                update_key="mapping:2",
                kind="mapping",
                reference_id="2",
                completed_at="2026-09-05T09:30:00+08:00",
                backup_path=second_backup,
                backup_sha256=second_hash,
                details={"batch_id": 2},
            )

        with self.assertRaisesRegex(ValueError, "全局最新"):
            rollback_latest_update(
                self.database,
                self.excel,
                first,
                rolled_back_at="2026-09-05T10:00:00+08:00",
            )

    def test_newer_untracked_mapping_update_blocks_older_tracked_update(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")
        with connect_database(self.database) as connection:
            connection.execute(
                """
                INSERT INTO mapping_import_batches (
                    source_file, source_sha256, total_rows, verified_rows,
                    pending_rows, inserted_rows, existing_rows, status, imported_at
                ) VALUES ('newer.xlsx', ?, 1, 0, 1, 0, 0, 'applied', ?)
                """,
                ("b" * 64, "2026-09-05T09:30:00+08:00"),
            )

        with self.assertRaisesRegex(ValueError, "全局最新"):
            rollback_latest_update(
                self.database,
                self.excel,
                key,
                rolled_back_at="2026-09-05T10:00:00+08:00",
            )

    def test_concurrent_writer_waits_until_rollback_commits(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")
        from cost_sync import update_rollback
        from cost_sync.locks import create_management_lock

        selected = threading.Event()
        continue_rollback = threading.Event()
        writer_finished = threading.Event()
        failures: list[BaseException] = []
        real_load = update_rollback._load_latest_record

        def paused_load(database: Path, update_key: str):
            record = real_load(database, update_key)
            selected.set()
            if not continue_rollback.wait(5):
                raise TimeoutError("test did not release rollback")
            return record

        def run_rollback() -> None:
            try:
                rollback_latest_update(
                    self.database,
                    self.excel,
                    key,
                    rolled_back_at="2026-09-05T10:00:00+08:00",
                )
            except BaseException as exc:
                failures.append(exc)

        def run_writer() -> None:
            try:
                with create_management_lock(self.database):
                    with connect_database(self.database) as connection:
                        connection.execute(
                            """
                            INSERT INTO mapping_import_batches (
                                source_file, source_sha256, total_rows, verified_rows,
                                pending_rows, inserted_rows, existing_rows, status, imported_at
                            ) VALUES ('concurrent.xlsx', ?, 1, 0, 1, 0, 0, 'applied', ?)
                            """,
                            ("c" * 64, "2026-09-05T10:30:00+08:00"),
                        )
            except BaseException as exc:
                failures.append(exc)
            finally:
                writer_finished.set()

        with mock.patch(
            "cost_sync.update_rollback._load_latest_record", side_effect=paused_load
        ):
            rollback_thread = threading.Thread(target=run_rollback)
            rollback_thread.start()
            self.assertTrue(selected.wait(5))
            writer_thread = threading.Thread(target=run_writer)
            writer_thread.start()
            self.assertFalse(writer_finished.wait(0.2))
            continue_rollback.set()
            rollback_thread.join(5)
            writer_thread.join(5)

        self.assertFalse(failures)
        self.assertTrue(writer_finished.is_set())
        with connect_database(self.database) as connection:
            self.assertIsNotNone(connection.execute(
                "SELECT 1 FROM mapping_import_batches WHERE source_file = 'concurrent.xlsx'"
            ).fetchone())

    def test_latest_mapping_update_restores_database_without_rebuilding_excel(self) -> None:
        backup = self.directory / "before-mapping.sqlite3.bak"
        backup_hash = backup_database(self.database, backup)
        with connect_database(self.database) as connection:
            batch_id = connection.execute(
                """
                INSERT INTO mapping_import_batches (
                    source_file, source_sha256, total_rows, verified_rows,
                    pending_rows, inserted_rows, existing_rows, status, imported_at
                ) VALUES ('shop.xlsx', ?, 1, 1, 0, 1, 0, 'applied', ?)
                """,
                ("a" * 64, "2026-09-05T09:00:00+08:00"),
            ).lastrowid
            connection.execute(
                """
                INSERT INTO douyin_sku_mappings (
                    douyin_sku_id, douyin_product_id, merchant_sku_code,
                    jushuitan_sku_id, product_name, specification, status,
                    import_batch_id, created_at, updated_at
                ) VALUES ('D-1', 'P-1', 'SKU-1', 'SKU-1', '商品', '规格',
                          'verified', ?, 'now', 'now')
                """,
                (batch_id,),
            )
            record_system_update(
                connection,
                update_key=f"mapping:{batch_id}",
                kind="mapping",
                reference_id=str(batch_id),
                completed_at="2026-09-05T09:00:00+08:00",
                backup_path=backup,
                backup_sha256=backup_hash,
                details={"batch_id": batch_id},
            )
        excel_before = hashlib.sha256(self.excel.read_bytes()).hexdigest()

        result = rollback_latest_update(
            self.database,
            self.excel,
            f"mapping:{batch_id}",
            rolled_back_at="2026-09-05T10:00:00+08:00",
        )

        with connect_database(self.database) as connection:
            mapping_count = connection.execute(
                "SELECT COUNT(*) FROM douyin_sku_mappings"
            ).fetchone()[0]
            status = connection.execute(
                "SELECT status FROM system_updates WHERE update_key = ?",
                (f"mapping:{batch_id}",),
            ).fetchone()[0]
        self.assertEqual(0, mapping_count)
        self.assertEqual("rolled_back", status)
        self.assertFalse(result.rebuilt_excel)
        self.assertEqual(excel_before, hashlib.sha256(self.excel.read_bytes()).hexdigest())

    def test_missing_or_hash_mismatched_backup_is_rejected(self) -> None:
        for damage in ("missing", "hash"):
            with self.subTest(damage=damage):
                with tempfile.TemporaryDirectory() as temp_dir:
                    database = Path(temp_dir) / "cost.sqlite3"
                    excel = Path(temp_dir) / "enabled.xlsx"
                    with connect_database(database) as connection:
                        upsert_products(
                            connection,
                            (Product("SKU", None, "商品", "1", 1, None, None),),
                            synced_at="2026-09-05T08:00:00+08:00",
                        )
                    backup = Path(temp_dir) / "before.sqlite3.bak"
                    backup_hash = backup_database(database, backup)
                    with connect_database(database) as connection:
                        record_system_update(
                            connection,
                            update_key="catalog:test",
                            kind="catalog",
                            reference_id="test",
                            completed_at="2026-09-05T09:00:00+08:00",
                            backup_path=backup,
                            backup_sha256=backup_hash,
                            details={"enabled_count": 1},
                        )
                    if damage == "missing":
                        backup.unlink()
                    else:
                        backup.write_bytes(backup.read_bytes() + b"changed")
                    with self.assertRaisesRegex(ValueError, "不存在|SHA-256"):
                        rollback_latest_update(
                            database,
                            excel,
                            "catalog:test",
                            rolled_back_at="2026-09-05T10:00:00+08:00",
                        )

    def test_partial_publish_failure_restores_current_database_and_excel(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")
        database_before = self._database_hash()
        excel_before = hashlib.sha256(self.excel.read_bytes()).hexdigest()
        real_replace = os.replace

        def fail_database_replace(source, target):
            if Path(target).resolve() == self.database.resolve():
                raise OSError("simulated database replace failure")
            return real_replace(source, target)

        with (
            mock.patch(
                "cost_sync.update_rollback.os.replace",
                side_effect=fail_database_replace,
            ),
            self.assertRaisesRegex(OSError, "simulated database replace failure"),
        ):
            rollback_latest_update(
                self.database,
                self.excel,
                key,
                rolled_back_at="2026-09-05T10:00:00+08:00",
            )

        self.assertEqual(database_before, self._database_hash())
        self.assertEqual(excel_before, hashlib.sha256(self.excel.read_bytes()).hexdigest())

    def test_snapshot_cleanup_failure_after_commit_still_reports_success(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")
        real_unlink = Path.unlink

        def fail_snapshot_unlink(path, *args, **kwargs):
            if path.suffix == ".previous":
                raise PermissionError("simulated locked snapshot")
            return real_unlink(path, *args, **kwargs)

        with mock.patch(
            "pathlib.Path.unlink", autospec=True, side_effect=fail_snapshot_unlink
        ):
            result = rollback_latest_update(
                self.database,
                self.excel,
                key,
                rolled_back_at="2026-09-05T10:00:00+08:00",
            )

        with connect_database(self.database) as connection:
            product = connection.execute(
                "SELECT cost_price FROM products WHERE sku_id = 'SKU-1'"
            ).fetchone()
            status = connection.execute(
                "SELECT status FROM system_updates WHERE update_key = ?", (key,)
            ).fetchone()[0]
        workbook = load_workbook(self.excel, read_only=True, data_only=True)
        try:
            excel_cost = workbook["启用商品成本价"]["C2"].value
        finally:
            workbook.close()

        self.assertEqual(key, result.update_key)
        self.assertEqual("1", product["cost_price"])
        self.assertEqual("rolled_back", status)
        self.assertEqual(1, excel_cost)

    def test_staged_integrity_failure_keeps_current_files(self) -> None:
        key = "catalog:2026-09-05T09:00:00+08:00"
        self._record_catalog_update(key, "2026-09-05T09:00:00+08:00")
        database_before = self._database_hash()
        excel_before = hashlib.sha256(self.excel.read_bytes()).hexdigest()
        from cost_sync import update_rollback

        real_check = update_rollback._require_integrity
        checks = 0

        def fail_staged_database_check(connection) -> None:
            nonlocal checks
            checks += 1
            if checks == 3:
                raise sqlite3.DatabaseError("simulated staged integrity failure")
            real_check(connection)

        with (
            mock.patch(
                "cost_sync.update_rollback._require_integrity",
                side_effect=fail_staged_database_check,
            ),
            self.assertRaisesRegex(sqlite3.DatabaseError, "simulated staged"),
        ):
            rollback_latest_update(
                self.database,
                self.excel,
                key,
                rolled_back_at="2026-09-05T10:00:00+08:00",
            )

        self.assertEqual(database_before, self._database_hash())
        self.assertEqual(excel_before, hashlib.sha256(self.excel.read_bytes()).hexdigest())

    def test_legacy_tracking_table_gains_recovery_and_details_columns(self) -> None:
        legacy = self.directory / "legacy.sqlite3"
        connection = sqlite3.connect(legacy)
        try:
            connection.execute(
                """
                CREATE TABLE system_updates (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    update_key TEXT NOT NULL UNIQUE,
                    kind TEXT NOT NULL,
                    reference_id TEXT NOT NULL,
                    status TEXT NOT NULL,
                    completed_at TEXT NOT NULL,
                    backup_file TEXT NOT NULL,
                    backup_sha256 TEXT NOT NULL
                )
                """
            )
            connection.commit()
        finally:
            connection.close()

        with connect_database(legacy) as connection:
            columns = {
                str(row[1]) for row in connection.execute(
                    "PRAGMA table_info(system_updates)"
                )
            }

        self.assertTrue(
            {"rolled_back_at", "recovery_backup_file", "recovery_backup_sha256", "details_json"}
            <= columns
        )


if __name__ == "__main__":
    unittest.main()
