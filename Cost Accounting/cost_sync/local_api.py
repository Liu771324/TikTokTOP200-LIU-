from __future__ import annotations

import argparse
import base64
import binascii
import hashlib
import json
import secrets
import sqlite3
import tempfile
import threading
import zipfile
from contextlib import closing, contextmanager
from dataclasses import asdict, dataclass
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse
from zoneinfo import ZoneInfo

from openpyxl.utils.exceptions import InvalidFileException

from .catalog_import import apply_catalog_import, preview_catalog_import
from .database import connect_database
from .locks import InterprocessMutex, create_management_lock
from .mapping_import import (
    apply_shop_mapping_import,
    normalize_shop_name,
    preview_mapping_import,
    preview_shop_mapping_changes,
)
from .mapping_review import capture_mapping_review_candidates
from .true_cost import estimate_true_cost


SHANGHAI = ZoneInfo("Asia/Shanghai")
EXTENSION_ORIGIN = "chrome-extension://kbdiohjlofljeafaddehaeciappaefka"
MANAGEMENT_ORIGIN = "app://cost-assistant-desktop"
MAX_ADMIN_BODY_BYTES = 48 * 1024 * 1024
MAX_CATALOG_FILE_BYTES = 16 * 1024 * 1024


def database_service_id(database_path: str | Path) -> str:
    """Return a non-reversible identifier for one local database location."""
    normalized = str(Path(database_path).resolve()).casefold().encode("utf-8")
    return hashlib.sha256(normalized).hexdigest()[:16]


class ProductRepository:
    def __init__(self, database_path: str | Path, *, environment: str = "unknown") -> None:
        self.database_path = Path(database_path).resolve()
        self.environment = environment

    def status(self) -> dict[str, object]:
        with closing(self._connect()) as connection:
            total_count = int(connection.execute("SELECT COUNT(*) FROM products").fetchone()[0])
            enabled_count = int(
                connection.execute(
                    "SELECT COUNT(*) FROM products WHERE enabled = 1"
                ).fetchone()[0]
            )
            verified_mapping_count = int(
                connection.execute(
                    "SELECT COUNT(*) FROM douyin_sku_mappings WHERE status = 'verified'"
                ).fetchone()[0]
            )
            row = connection.execute(
                "SELECT value FROM sync_state WHERE key = 'last_successful_sync'"
            ).fetchone()
            mapping_shops = [
                str(item[0])
                for item in connection.execute(
                    "SELECT DISTINCT shop_name FROM douyin_sku_mappings "
                    "WHERE shop_name <> '历史导入' ORDER BY shop_name"
                )
            ]
        return {
            "status": "ok",
            "environment": self.environment,
            "service_id": database_service_id(self.database_path),
            "total_count": total_count,
            "enabled_count": enabled_count,
            "verified_mapping_count": verified_mapping_count,
            "mapping_shops": mapping_shops,
            "capabilities": ["true_cost_breakdown_v1", "page_category_context_v1"],
            "last_successful_sync": None if row is None else row[0],
        }

    def get_enabled_product(self, sku_id: str) -> dict[str, object] | None:
        with closing(self._connect()) as connection:
            row = connection.execute(
                """
                SELECT sku_id, name, cost_price, synced_at
                FROM products
                WHERE sku_id = ? AND enabled = 1
                """,
                (sku_id,),
            ).fetchone()
        return None if row is None else dict(row)

    def get_enabled_products(self, sku_ids: list[str]) -> list[dict[str, object]]:
        if not sku_ids:
            return []
        if len(sku_ids) > 100:
            raise ValueError("单次最多查询 100 个商品编码")
        placeholders = ",".join("?" for _ in sku_ids)
        with closing(self._connect()) as connection:
            rows = connection.execute(
                f"""
                SELECT sku_id, name, cost_price, synced_at
                FROM products
                WHERE enabled = 1 AND sku_id IN ({placeholders})
                """,
                sku_ids,
            ).fetchall()
        by_sku = {row["sku_id"]: dict(row) for row in rows}
        return [by_sku[sku_id] for sku_id in sku_ids if sku_id in by_sku]

    def get_verified_mappings(
        self, douyin_sku_ids: list[str]
    ) -> list[dict[str, object]]:
        if not douyin_sku_ids:
            return []
        if len(douyin_sku_ids) > 100:
            raise ValueError("单次最多查询 100 个抖店 SKUID")
        placeholders = ",".join("?" for _ in douyin_sku_ids)
        with closing(self._connect()) as connection:
            rows = connection.execute(
                f"""
                SELECT douyin_sku_id, jushuitan_sku_id, status, import_batch_id
                FROM douyin_sku_mappings
                WHERE status = 'verified' AND douyin_sku_id IN ({placeholders})
                """,
                douyin_sku_ids,
            ).fetchall()
        by_sku = {row["douyin_sku_id"]: dict(row) for row in rows}
        return [
            by_sku[douyin_sku_id]
            for douyin_sku_id in douyin_sku_ids
            if douyin_sku_id in by_sku
        ]

    def get_true_costs(
        self,
        douyin_sku_ids: list[str],
        lowest_prices: list[str],
        *,
        region: str = "",
        page_category_levels: list[tuple[str, str, str, str] | None] | None = None,
    ) -> list[dict[str, object]]:
        if len(douyin_sku_ids) != len(lowest_prices):
            raise ValueError("抖店 SKUID 与最低到手价必须一一对应")
        if len(douyin_sku_ids) > 100:
            raise ValueError("单次最多查询 100 个抖店 SKUID")
        if page_category_levels is not None and len(page_category_levels) != len(
            douyin_sku_ids
        ):
            raise ValueError("抖店 SKUID 与页面类目必须一一对应")
        effective_on = datetime.now(SHANGHAI).date().isoformat()
        return [
            asdict(
                estimate_true_cost(
                    self.database_path,
                    douyin_sku_id=douyin_sku_id,
                    lowest_price=lowest_price,
                    page_category_levels=page_categories,
                    effective_on=effective_on,
                    region=region,
                )
            )
            for douyin_sku_id, lowest_price, page_categories in zip(
                douyin_sku_ids,
                lowest_prices,
                page_category_levels or [None] * len(douyin_sku_ids),
                strict=True,
            )
        ]

    def _connect(self) -> sqlite3.Connection:
        if not self.database_path.is_file():
            raise FileNotFoundError(f"SQLite 文件不存在：{self.database_path}")
        uri = f"{self.database_path.as_uri()}?mode=ro"
        connection = sqlite3.connect(uri, uri=True)
        connection.row_factory = sqlite3.Row
        return connection


@dataclass(frozen=True)
class CatalogPreviewSession:
    preview_id: str
    ordinary_name: str
    combination_name: str
    ordinary_bytes: bytes
    combination_bytes: bytes
    ordinary_sha256: str
    combination_sha256: str
    database_revision: tuple[int, int]


class CatalogImportManager:
    def __init__(
        self, database_path: str | Path, output_excel_path: str | Path, *, lock=None
    ) -> None:
        self.database_path = Path(database_path).resolve()
        self.output_excel_path = Path(output_excel_path).resolve()
        self._lock = lock or threading.Lock()
        self._session: CatalogPreviewSession | None = None

    def preview(self, payload: object) -> dict[str, object]:
        if not isinstance(payload, dict):
            raise ValueError("预览请求格式无效")
        ordinary_name, ordinary_bytes = _decode_catalog_upload(
            payload.get("ordinary"), "普通商品表"
        )
        combination_name, combination_bytes = _decode_catalog_upload(
            payload.get("combination"), "组合商品表"
        )
        if len(ordinary_bytes) + len(combination_bytes) > MAX_CATALOG_FILE_BYTES * 2:
            raise ValueError("两份商品表合计大小超过安全上限")

        with self._lock:
            with self._materialize(ordinary_bytes, combination_bytes) as paths:
                preview = preview_catalog_import(*paths)
                changes = self._compare_with_current(preview.products)
            preview_id = secrets.token_urlsafe(24)
            self._session = CatalogPreviewSession(
                preview_id=preview_id,
                ordinary_name=ordinary_name,
                combination_name=combination_name,
                ordinary_bytes=ordinary_bytes,
                combination_bytes=combination_bytes,
                ordinary_sha256=preview.ordinary_sha256,
                combination_sha256=preview.combination_sha256,
                database_revision=_database_revision(self.database_path),
            )
            return {
                "preview_id": preview_id,
                "ready_to_import": preview.ready_to_import,
                "ordinary_file": ordinary_name,
                "combination_file": combination_name,
                "ordinary_rows": preview.ordinary_rows,
                "ordinary_products": preview.ordinary_products,
                "combination_rows": preview.combination_rows,
                "combination_products": preview.combination_products,
                "merged_products": preview.merged_products,
                "empty_cost_products": preview.empty_cost_products,
                "resolved_overlap_count": len(preview.resolved_overlaps),
                "conflicting_overlaps": list(preview.conflicting_overlaps),
                "combination_cost_mismatches": list(preview.combination_cost_mismatches),
                **changes,
            }

    def apply(self, payload: object) -> dict[str, object]:
        if not isinstance(payload, dict) or not isinstance(payload.get("preview_id"), str):
            raise ValueError("确认请求缺少有效预览凭证")
        with self._lock:
            session = self._session
            if session is None or not secrets.compare_digest(
                session.preview_id, payload["preview_id"]
            ):
                raise ValueError("预览已失效，请重新选择两份商品表并预览")
            if _database_revision(self.database_path) != session.database_revision:
                self._session = None
                raise ValueError("成本库在预览后已由其他窗口更新，请重新预览")
            with self._materialize(
                session.ordinary_bytes, session.combination_bytes
            ) as paths:
                preview = preview_catalog_import(*paths)
                result = apply_catalog_import(
                    self.database_path,
                    self.output_excel_path,
                    preview,
                    confirmed_ordinary_sha256=session.ordinary_sha256,
                    confirmed_combination_sha256=session.combination_sha256,
                    imported_at=datetime.now(SHANGHAI).replace(microsecond=0).isoformat(),
                )
            self._session = None
            return {
                "status": "applied",
                "total_count": result.total_count,
                "enabled_count": result.enabled_count,
                "empty_cost_products": result.empty_cost_products,
                "completed_at": result.completed_at,
                "backup_file": None
                if result.backup_path is None
                else result.backup_path.name,
            }

    @contextmanager
    def _materialize(self, ordinary: bytes, combination: bytes):
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            prefix=".catalog-import-", dir=self.database_path.parent
        ) as temp_dir:
            root = Path(temp_dir)
            ordinary_path = root / "ordinary.xlsx"
            combination_path = root / "combination.xlsx"
            ordinary_path.write_bytes(ordinary)
            combination_path.write_bytes(combination)
            yield ordinary_path, combination_path

    def _compare_with_current(self, products) -> dict[str, int]:
        with closing(sqlite3.connect(
            f"{self.database_path.as_uri()}?mode=ro", uri=True
        )) as connection:
            current = {
                str(row[0]): (row[1], int(row[2]))
                for row in connection.execute(
                    "SELECT sku_id, cost_price, enabled FROM products"
                )
            }
        incoming = {
            product.sku_id: (product.cost_price, product.enabled) for product in products
        }
        shared = current.keys() & incoming.keys()
        return {
            "added_products": len(incoming.keys() - current.keys()),
            "removed_products": len(current.keys() - incoming.keys()),
            "cost_changed_products": sum(
                current[sku_id][0] != incoming[sku_id][0] for sku_id in shared
            ),
            "status_changed_products": sum(
                current[sku_id][1] != incoming[sku_id][1] for sku_id in shared
            ),
        }


@dataclass(frozen=True)
class MappingPreviewSession:
    preview_id: str
    source_name: str
    source_bytes: bytes
    source_sha256: str
    shop_name: str
    ready_to_import: bool
    database_revision: tuple[int, int]


class MappingImportManager:
    def __init__(self, database_path: str | Path, *, lock=None) -> None:
        self.database_path = Path(database_path).resolve()
        self._lock = lock or threading.Lock()
        self._session: MappingPreviewSession | None = None

    def preview(self, payload: object) -> dict[str, object]:
        if not isinstance(payload, dict):
            raise ValueError("预览请求格式无效")
        source_name, source_bytes = _decode_catalog_upload(
            payload.get("mapping"), "抖店商品表"
        )
        shop_name = normalize_shop_name(payload.get("shop_name"))
        with self._lock:
            with self._materialize(source_bytes) as source_path:
                preview = preview_mapping_import(self.database_path, source_path)
                changes = preview_shop_mapping_changes(
                    self.database_path, preview, shop_name
                )
            ready = (
                preview.ready_to_import
                and not changes.conflicting_sku_ids
                and not (changes.current_shop_rows > 0 and changes.shared_shop_rows == 0)
                and any((
                    changes.new_rows,
                    changes.changed_rows,
                    changes.removed_rows,
                    changes.legacy_rows,
                ))
            )
            already_current = (
                preview.ready_to_import
                and not changes.conflicting_sku_ids
                and not (changes.current_shop_rows > 0 and changes.shared_shop_rows == 0)
                and not any((
                    changes.new_rows,
                    changes.changed_rows,
                    changes.removed_rows,
                    changes.legacy_rows,
                ))
                and preview.verified_rows > 0
            )
            preview_id = secrets.token_urlsafe(24)
            self._session = MappingPreviewSession(
                preview_id=preview_id,
                source_name=source_name,
                source_bytes=source_bytes,
                source_sha256=preview.source_sha256,
                shop_name=shop_name,
                ready_to_import=ready,
                database_revision=_database_revision(self.database_path),
            )
            return {
                "preview_id": preview_id,
                "ready_to_import": ready,
                "already_current": already_current,
                "shop_name": shop_name,
                "source_file": source_name,
                "total_rows": preview.total_rows,
                "unique_douyin_skus": preview.unique_douyin_skus,
                "matched_rows": preview.verified_rows,
                "pending_rows": preview.pending_rows,
                "duplicate_rows": preview.duplicate_rows,
                "invalid_rows": preview.invalid_rows,
                "complete_category_rows": preview.complete_category_rows,
                "incomplete_category_rows": preview.incomplete_category_rows,
                "service_fee_matched_rows": preview.service_fee_matched_rows,
                "service_fee_unconfigured_rows": preview.service_fee_unconfigured_rows,
                "service_fee_disabled_rows": preview.service_fee_disabled_rows,
                "major_category_configured_rows": preview.major_category_configured_rows,
                "major_category_unconfigured_rows": preview.major_category_unconfigured_rows,
                "new_rows": changes.new_rows,
                "existing_rows": changes.existing_rows,
                "changed_rows": changes.changed_rows,
                "removed_rows": changes.removed_rows,
                "legacy_rows": changes.legacy_rows,
                "current_shop_rows": changes.current_shop_rows,
                "shared_shop_rows": changes.shared_shop_rows,
                "shop_file_mismatch": (
                    changes.current_shop_rows > 0 and changes.shared_shop_rows == 0
                ),
                "conflict_count": len(changes.conflicting_sku_ids),
                "conflicting_sku_ids": list(changes.conflicting_sku_ids[:20]),
                "review_count": preview.pending_rows + len(changes.conflicting_sku_ids),
                "can_save_reviews": (
                    preview.ready_to_import
                    and not (changes.current_shop_rows > 0 and changes.shared_shop_rows == 0)
                    and (preview.pending_rows > 0 or bool(changes.conflicting_sku_ids))
                ),
            }

    def apply(self, payload: object) -> dict[str, object]:
        if not isinstance(payload, dict) or not isinstance(payload.get("preview_id"), str):
            raise ValueError("确认请求缺少有效预览凭证")
        with self._lock:
            session = self._session
            if session is None or not secrets.compare_digest(
                session.preview_id, payload["preview_id"]
            ):
                raise ValueError("预览已失效，请重新选择抖店商品表并预览")
            if _database_revision(self.database_path) != session.database_revision:
                self._session = None
                raise ValueError("成本库或店铺映射在预览后已由其他窗口更新，请重新预览")
            if not session.ready_to_import:
                raise ValueError("本次预览没有可导入的新映射")
            with self._materialize(session.source_bytes) as source_path:
                preview = preview_mapping_import(self.database_path, source_path)
                changes = preview_shop_mapping_changes(
                    self.database_path, preview, session.shop_name
                )
                if not secrets.compare_digest(preview.source_sha256, session.source_sha256):
                    raise ValueError("抖店商品表内容已变化，请重新预览")
                if (
                    not preview.ready_to_import
                    or changes.conflicting_sku_ids
                    or (changes.current_shop_rows > 0 and changes.shared_shop_rows == 0)
                    or not any((
                        changes.new_rows,
                        changes.changed_rows,
                        changes.removed_rows,
                        changes.legacy_rows,
                    ))
                ):
                    raise ValueError("确认前复核未通过，请重新预览抖店商品表")
                result = apply_shop_mapping_import(
                    self.database_path,
                    preview,
                    shop_name=session.shop_name,
                    imported_at=datetime.now(SHANGHAI).replace(microsecond=0).isoformat(),
                )
            self._session = None
            return {
                "status": "applied",
                "batch_id": result.batch_id,
                "inserted_rows": result.inserted_rows,
                "existing_rows": result.existing_rows,
                "updated_rows": result.updated_rows,
                "removed_rows": result.removed_rows,
                "legacy_rows": result.legacy_rows,
                "shop_name": session.shop_name,
                "backup_file": result.backup_path.name,
            }

    def save_reviews(self, payload: object) -> dict[str, object]:
        if not isinstance(payload, dict) or not isinstance(payload.get("preview_id"), str):
            raise ValueError("保存请求缺少有效预览凭证")
        with self._lock:
            session = self._session
            if session is None or not secrets.compare_digest(
                session.preview_id, payload["preview_id"]
            ):
                raise ValueError("预览已失效，请重新选择抖店商品表并预览")
            if _database_revision(self.database_path) != session.database_revision:
                self._session = None
                raise ValueError("成本库或店铺映射在预览后已更新，请重新预览")
            with self._materialize(session.source_bytes) as source_path:
                preview = preview_mapping_import(self.database_path, source_path)
                changes = preview_shop_mapping_changes(
                    self.database_path, preview, session.shop_name
                )
                if (
                    not preview.ready_to_import
                    or (changes.current_shop_rows > 0 and changes.shared_shop_rows == 0)
                ):
                    raise ValueError("当前预览不能保存待处理项，请先核对文件和店铺")
                result = capture_mapping_review_candidates(
                    self.database_path,
                    preview,
                    shop_name=session.shop_name,
                    conflict_sku_ids=changes.conflicting_sku_ids,
                    captured_at=datetime.now(SHANGHAI).replace(microsecond=0).isoformat(),
                )
            self._session = None
            return {
                "status": "saved",
                "batch_id": result.batch_id,
                "saved_rows": result.saved_rows,
                "unchanged_rows": result.unchanged_rows,
                "shop_name": session.shop_name,
                "backup_file": None if result.backup_path is None else result.backup_path.name,
            }

    @contextmanager
    def _materialize(self, source: bytes):
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            prefix=".mapping-import-", dir=self.database_path.parent
        ) as temp_dir:
            source_path = Path(temp_dir) / "douyin-mappings.xlsx"
            source_path.write_bytes(source)
            yield source_path

class LocalApiServer(ThreadingHTTPServer):
    daemon_threads = True


def create_server(
    database_path: str | Path,
    *,
    port: int = 8765,
    allowed_origins: tuple[str, ...] = (),
    environment: str = "unknown",
    output_excel_path: str | Path | None = None,
    catalog_manager: CatalogImportManager | None = None,
    mapping_manager: MappingImportManager | None = None,
) -> LocalApiServer:
    if "*" in allowed_origins:
        raise ValueError("不允许使用通配符跨域来源")
    if environment not in {"test", "production", "unknown"}:
        raise ValueError("environment 必须是 test、production 或 unknown")
    repository = ProductRepository(database_path, environment=environment)
    with connect_database(database_path) as connection:
        connection.commit()
    repository.status()
    database = Path(database_path).resolve()
    output_excel = (
        database.parent / "enabled_product_costs.xlsx"
        if output_excel_path is None
        else Path(output_excel_path).resolve()
    )
    if (catalog_manager is None) != (mapping_manager is None):
        raise ValueError("商品与映射管理器必须同时提供")
    if catalog_manager is None:
        management_lock = create_management_lock(database)
        catalog_manager = CatalogImportManager(database, output_excel, lock=management_lock)
        mapping_manager = MappingImportManager(database, lock=management_lock)
    handler = _make_handler(
        repository, catalog_manager, mapping_manager, frozenset(allowed_origins)
    )
    return LocalApiServer(("127.0.0.1", port), handler)


def _make_handler(
    repository: ProductRepository,
    catalog_manager: CatalogImportManager,
    mapping_manager: MappingImportManager,
    allowed_origins: frozenset[str],
):
    class Handler(BaseHTTPRequestHandler):
        server_version = "CostAccountingLocal/1"

        def do_GET(self) -> None:
            try:
                origin = self.headers.get("Origin")
                if origin is not None and origin not in allowed_origins:
                    self._send_error(HTTPStatus.FORBIDDEN, "跨域来源未获允许")
                    return
                parsed = urlparse(self.path)
                if parsed.path == "/health":
                    self._send_json(HTTPStatus.OK, repository.status())
                    return
                if parsed.path == "/api/v1/products":
                    query = parse_qs(parsed.query)
                    sku_ids = _parse_repeated_or_legacy_sku_ids(
                        query.get("sku_id", []), query.get("sku_ids", [])
                    )
                    self._send_json(
                        HTTPStatus.OK,
                        {"products": repository.get_enabled_products(sku_ids)},
                    )
                    return
                if parsed.path == "/api/v1/mappings":
                    query = parse_qs(parsed.query)
                    douyin_sku_ids = _parse_repeated_or_legacy_sku_ids(
                        query.get("douyin_sku_id", []), query.get("douyin_sku_ids", [])
                    )
                    self._send_json(
                        HTTPStatus.OK,
                        {"mappings": repository.get_verified_mappings(douyin_sku_ids)},
                    )
                    return
                if parsed.path == "/api/v1/true-costs":
                    query = parse_qs(parsed.query, keep_blank_values=True)
                    douyin_sku_ids = [
                        value.strip()
                        for value in query.get("douyin_sku_id", [])
                        if value.strip()
                    ]
                    lowest_prices = [
                        value.strip()
                        for value in query.get("lowest_price", [])
                        if value.strip()
                    ]
                    category_columns = [
                        [value.strip() for value in query.get(f"category_level_{index}", [])]
                        for index in range(1, 5)
                    ]
                    page_categories = None
                    if any(category_columns):
                        if any(
                            len(column) != len(douyin_sku_ids)
                            for column in category_columns
                        ):
                            raise ValueError("抖店 SKUID 与页面类目必须一一对应")
                        page_categories = []
                        for values in zip(*category_columns, strict=True):
                            if all(not value for value in values):
                                page_categories.append(None)
                            elif all(values):
                                page_categories.append(tuple(values))  # type: ignore[arg-type]
                            else:
                                raise ValueError("页面类目必须提供完整四级类目")
                    regions = query.get("region", [])
                    if len(regions) > 1:
                        raise ValueError("收货地区只能指定一次")
                    self._send_json(
                        HTTPStatus.OK,
                        {
                            "true_costs": repository.get_true_costs(
                                douyin_sku_ids,
                                lowest_prices,
                                region=regions[0].strip() if regions else "",
                                page_category_levels=page_categories,
                            )
                        },
                    )
                    return
                prefix = "/api/v1/products/"
                if parsed.path.startswith(prefix):
                    sku_id = unquote(parsed.path[len(prefix) :]).strip()
                    if not sku_id:
                        self._send_error(HTTPStatus.BAD_REQUEST, "商品编码不能为空")
                        return
                    product = repository.get_enabled_product(sku_id)
                    if product is None:
                        self._send_error(HTTPStatus.NOT_FOUND, "未找到启用商品")
                        return
                    self._send_json(HTTPStatus.OK, product)
                    return
                self._send_error(HTTPStatus.NOT_FOUND, "接口不存在")
            except ValueError as exc:
                self._send_error(HTTPStatus.BAD_REQUEST, str(exc))
            except (OSError, sqlite3.Error):
                self._send_error(HTTPStatus.SERVICE_UNAVAILABLE, "本地成本数据库暂不可用")

        def do_OPTIONS(self) -> None:
            origin = self.headers.get("Origin")
            if origin not in allowed_origins:
                self._send_error(HTTPStatus.FORBIDDEN, "跨域来源未获允许")
                return
            parsed = urlparse(self.path)
            if parsed.path.startswith("/api/v1/admin/") and origin != MANAGEMENT_ORIGIN:
                self._send_error(HTTPStatus.FORBIDDEN, "管理操作来源未获允许")
                return
            self.send_response(HTTPStatus.NO_CONTENT)
            self._send_common_headers(origin)
            methods = "GET, OPTIONS"
            if origin == MANAGEMENT_ORIGIN:
                methods = "GET, POST, OPTIONS"
            self.send_header("Access-Control-Allow-Methods", methods)
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Access-Control-Max-Age", "600")
            self.end_headers()

        def do_POST(self) -> None:
            origin = self.headers.get("Origin")
            if origin != MANAGEMENT_ORIGIN or origin not in allowed_origins:
                self._send_error(HTTPStatus.FORBIDDEN, "管理操作来源未获允许")
                return
            parsed = urlparse(self.path)
            try:
                if parsed.path == "/api/v1/admin/catalog/preview":
                    self._send_json(
                        HTTPStatus.OK, catalog_manager.preview(self._read_json())
                    )
                    return
                if parsed.path == "/api/v1/admin/catalog/import":
                    self._send_json(
                        HTTPStatus.OK, catalog_manager.apply(self._read_json())
                    )
                    return
                if parsed.path == "/api/v1/admin/mappings/preview":
                    self._send_json(
                        HTTPStatus.OK, mapping_manager.preview(self._read_json())
                    )
                    return
                if parsed.path == "/api/v1/admin/mappings/import":
                    self._send_json(
                        HTTPStatus.OK, mapping_manager.apply(self._read_json())
                    )
                    return
                self._send_error(HTTPStatus.METHOD_NOT_ALLOWED, "查询接口只允许读取")
            except (ValueError, json.JSONDecodeError, binascii.Error) as exc:
                self._send_error(HTTPStatus.BAD_REQUEST, str(exc))
            except (zipfile.BadZipFile, InvalidFileException):
                self._send_error(HTTPStatus.BAD_REQUEST, "上传文件不是有效的 XLSX 文件")
            except (OSError, sqlite3.Error):
                self._send_error(HTTPStatus.SERVICE_UNAVAILABLE, "本机数据更新暂时无法完成")

        def log_message(self, format: str, *args: object) -> None:
            return

        def _send_error(self, status: HTTPStatus, message: str) -> None:
            self._send_json(status, {"error": message})

        def _send_json(self, status: HTTPStatus, payload: object) -> None:
            body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode(
                "utf-8"
            )
            self.send_response(status)
            self._send_common_headers(self.headers.get("Origin"))
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def _send_common_headers(self, origin: str | None) -> None:
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            if origin in allowed_origins:
                self.send_header("Access-Control-Allow-Origin", origin)
                self.send_header("Vary", "Origin")

        def _read_json(self) -> object:
            content_type = self.headers.get("Content-Type", "").split(";", 1)[0].strip()
            if content_type != "application/json":
                raise ValueError("管理请求必须使用 JSON 格式")
            raw_length = self.headers.get("Content-Length")
            if raw_length is None:
                raise ValueError("管理请求缺少内容长度")
            try:
                length = int(raw_length)
            except ValueError as exc:
                raise ValueError("管理请求内容长度无效") from exc
            if length <= 0 or length > MAX_ADMIN_BODY_BYTES:
                raise ValueError("管理请求大小超过安全上限")
            return json.loads(self.rfile.read(length).decode("utf-8"))

    return Handler


def _database_revision(database_path: str | Path) -> tuple[int, int]:
    stat = Path(database_path).resolve().stat()
    return stat.st_size, stat.st_mtime_ns


def _decode_catalog_upload(value: object, label: str) -> tuple[str, bytes]:
    if not isinstance(value, dict):
        raise ValueError(f"请选择{label}")
    name = value.get("name")
    encoded = value.get("content_base64")
    if (
        not isinstance(name, str)
        or not name.strip()
        or Path(name).name != name
        or Path(name).suffix.lower() != ".xlsx"
    ):
        raise ValueError(f"{label}必须是 XLSX 文件")
    if not isinstance(encoded, str) or not encoded:
        raise ValueError(f"{label}内容为空")
    content = base64.b64decode(encoded, validate=True)
    if not content or len(content) > MAX_CATALOG_FILE_BYTES:
        raise ValueError(f"{label}大小超过安全上限")
    return name, content


def _parse_sku_ids(value: str) -> list[str]:
    sku_ids = list(dict.fromkeys(part.strip() for part in value.split(",") if part.strip()))
    if len(sku_ids) > 100:
        raise ValueError("单次最多查询 100 个商品编码")
    return sku_ids


def _parse_repeated_or_legacy_sku_ids(
    repeated_values: list[str], legacy_values: list[str]
) -> list[str]:
    if repeated_values:
        sku_ids = list(dict.fromkeys(value.strip() for value in repeated_values if value.strip()))
        if len(sku_ids) > 100:
            raise ValueError("单次最多查询 100 个商品编码")
        return sku_ids
    return _parse_sku_ids(legacy_values[0] if legacy_values else "")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="启动成本助手本机接口")
    parser.add_argument("--database", required=True, help="SQLite 数据库路径")
    parser.add_argument("--port", type=int, default=8765, help="本机端口，默认 8765")
    parser.add_argument(
        "--environment",
        choices=("test", "production"),
        default="test",
        help="成本数据环境，默认 test；正式数据必须显式指定 production",
    )
    parser.add_argument(
        "--allowed-origin",
        action="append",
        default=[],
        help="允许跨域访问的浏览器扩展来源，可重复填写；默认不允许",
    )
    parser.add_argument(
        "--output-excel",
        help="启用商品成本 Excel 路径；默认与数据库同目录",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if not 1 <= args.port <= 65535:
        print("启动失败：port 必须在 1 到 65535 之间")
        return 1
    try:
        server = create_server(
            args.database,
            port=args.port,
            allowed_origins=tuple(args.allowed_origin),
            environment=args.environment,
            output_excel_path=args.output_excel,
        )
    except (OSError, sqlite3.Error, ValueError) as exc:
        print(f"启动失败：{exc}")
        return 1

    print(f"成本助手本机接口：http://127.0.0.1:{args.port}")
    print("按 Ctrl+C 停止。")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
