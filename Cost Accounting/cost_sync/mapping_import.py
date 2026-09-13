from __future__ import annotations

import argparse
import hashlib
import sqlite3
from collections import Counter
from contextlib import closing
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .database import connect_database
from .locks import management_locked
from .mapping_review import (
    mark_review_candidates_resolved_by_import,
    upsert_mapping_review_candidates,
)
from .service_fee import score_service_fee_category_rule
from .shipping import parse_specification_weight
from .update_rollback import backup_database, record_system_update
from .xlsx_safety import MAX_IMPORT_ROWS, validate_xlsx_container


REQUIRED_HEADERS = (
    "商品ID",
    "商品名称",
    "一级类目",
    "二级类目",
    "三级类目",
    "四级类目",
    "商家SKU编码",
    "规格ID（SKUID）",
    "商品规格",
)


@dataclass(frozen=True)
class MappingCandidate:
    source_row: int
    douyin_sku_id: str
    douyin_product_id: str
    merchant_sku_code: str
    product_name: str
    specification: str
    category_level_1: str
    category_level_2: str
    category_level_3: str
    category_level_4: str
    status: str
    jushuitan_sku_id: str | None
    weight_kg: str | None = None
    weight_source: str = ""
    weight_unconfigured_reason: str = ""
    major_category: str = ""
    major_category_unconfigured_reason: str = ""


@dataclass(frozen=True)
class MappingImportPreview:
    source_file: str
    source_sha256: str
    total_rows: int
    unique_douyin_skus: int
    verified_rows: int
    pending_rows: int
    duplicate_rows: int
    invalid_rows: int
    candidates: tuple[MappingCandidate, ...]
    complete_category_rows: int = 0
    incomplete_category_rows: int = 0
    service_fee_matched_rows: int = 0
    service_fee_unconfigured_rows: int = 0
    service_fee_disabled_rows: int = 0
    major_category_configured_rows: int = 0
    major_category_unconfigured_rows: int = 0

    @property
    def ready_to_import(self) -> bool:
        return self.duplicate_rows == 0 and self.invalid_rows == 0


@dataclass(frozen=True)
class MappingImportResult:
    batch_id: int
    inserted_rows: int
    existing_rows: int
    reused_batch: bool


@dataclass(frozen=True)
class ShopMappingChanges:
    new_rows: int
    existing_rows: int
    changed_rows: int
    removed_rows: int
    legacy_rows: int
    current_shop_rows: int
    shared_shop_rows: int
    conflicting_sku_ids: tuple[str, ...]


@dataclass(frozen=True)
class ShopMappingImportResult:
    batch_id: int
    inserted_rows: int
    existing_rows: int
    updated_rows: int
    removed_rows: int
    legacy_rows: int
    backup_path: Path


def preview_mapping_import(
    database_path: str | Path, source_path: str | Path
) -> MappingImportPreview:
    database = Path(database_path).resolve()
    source = Path(source_path).resolve()
    if not database.is_file():
        raise FileNotFoundError(f"SQLite 文件不存在：{database}")
    if not source.is_file():
        raise FileNotFoundError(f"抖店导出文件不存在：{source}")
    validate_xlsx_container(source, "抖店商品表")

    enabled_sku_ids = _load_enabled_sku_ids(database)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("抖店导出文件必须且只能包含一个工作表")
        worksheet = workbook[workbook.sheetnames[0]]
        # 抖店导出文件可能把实际 21 列、数万行错误声明为 A1:A1。
        # 只读模式默认信任该元数据，重置后才会扫描真实工作表范围。
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError("抖店导出文件为空") from exc
        headers = tuple(_header_text(value) for value in raw_headers)
        duplicate_headers = [
            header for header, count in Counter(headers).items() if header and count > 1
        ]
        if duplicate_headers:
            raise ValueError(f"抖店导出文件存在重复列：{', '.join(duplicate_headers)}")
        missing_headers = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing_headers:
            raise ValueError(f"抖店导出文件缺少必填列：{', '.join(missing_headers)}")
        indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}

        candidates: list[MappingCandidate] = []
        seen_skus: set[str] = set()
        duplicate_rows = 0
        invalid_rows = 0
        total_rows = 0
        verified_rows = 0
        pending_rows = 0

        for source_row, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            total_rows += 1
            if total_rows > MAX_IMPORT_ROWS:
                raise ValueError(f"抖店商品表数据行超过 {MAX_IMPORT_ROWS:,} 行安全上限")
            try:
                douyin_sku_id = _identifier(
                    row[indexes["规格ID（SKUID）"]], "规格ID（SKUID）", source_row
                )
                douyin_product_id = _identifier(
                    row[indexes["商品ID"]], "商品ID", source_row
                )
                merchant_sku_code = _identifier(
                    row[indexes["商家SKU编码"]], "商家SKU编码", source_row
                )
            except ValueError:
                invalid_rows += 1
                continue

            if douyin_sku_id in seen_skus:
                duplicate_rows += 1
                continue
            seen_skus.add(douyin_sku_id)

            if merchant_sku_code in enabled_sku_ids:
                status = "verified"
                jushuitan_sku_id = merchant_sku_code
                verified_rows += 1
            else:
                status = "pending"
                jushuitan_sku_id = None
                pending_rows += 1
            specification = _display_text(row[indexes["商品规格"]])
            parsed_weight = parse_specification_weight(specification)
            candidates.append(
                MappingCandidate(
                    source_row=source_row,
                    douyin_sku_id=douyin_sku_id,
                    douyin_product_id=douyin_product_id,
                    merchant_sku_code=merchant_sku_code,
                    product_name=_display_text(row[indexes["商品名称"]]),
                    specification=specification,
                    category_level_1=_category_text(row[indexes["一级类目"]]),
                    category_level_2=_category_text(row[indexes["二级类目"]]),
                    category_level_3=_category_text(row[indexes["三级类目"]]),
                    category_level_4=_category_text(row[indexes["四级类目"]]),
                    status=status,
                    jushuitan_sku_id=jushuitan_sku_id,
                    weight_kg=parsed_weight.weight_kg,
                    weight_source=parsed_weight.source or "",
                    weight_unconfigured_reason=parsed_weight.unconfigured_reason or "",
                )
            )
    finally:
        workbook.close()

    candidates = _enrich_major_categories(database, candidates)
    category_coverage = _summarize_service_fee_coverage(database, candidates)
    return MappingImportPreview(
        source_file=source.name,
        source_sha256=_sha256(source),
        total_rows=total_rows,
        unique_douyin_skus=len(seen_skus),
        verified_rows=verified_rows,
        pending_rows=pending_rows,
        duplicate_rows=duplicate_rows,
        invalid_rows=invalid_rows,
        candidates=tuple(candidates),
        **category_coverage,
    )


def _summarize_service_fee_coverage(
    database_path: Path, candidates: list[MappingCandidate]
) -> dict[str, int]:
    effective_on = datetime.now().astimezone().date().isoformat()
    with connect_database(database_path) as connection:
        rows = connection.execute(
            """
            SELECT id, major_category, category_level_1, category_level_2,
                   category_level_3, category_level_4, enabled, effective_from
            FROM basic_service_fee_rates
            WHERE archived = 0 AND effective_from <= ?
            ORDER BY
              ((category_level_1 <> '/') + (category_level_2 <> '/') +
               (category_level_3 <> '/') + (category_level_4 <> '/')) DESC,
              effective_from DESC, id DESC
            """,
            (effective_on,),
        ).fetchall()
    counts = {
        "complete_category_rows": 0,
        "incomplete_category_rows": 0,
        "service_fee_matched_rows": 0,
        "service_fee_unconfigured_rows": 0,
        "service_fee_disabled_rows": 0,
        "major_category_configured_rows": 0,
        "major_category_unconfigured_rows": 0,
    }
    for candidate in candidates:
        if candidate.status != "verified":
            continue
        if candidate.major_category:
            counts["major_category_configured_rows"] += 1
        else:
            counts["major_category_unconfigured_rows"] += 1
        category = (
            candidate.category_level_1,
            candidate.category_level_2,
            candidate.category_level_3,
            candidate.category_level_4,
        )
        if any(not value for value in category):
            counts["incomplete_category_rows"] += 1
            continue
        counts["complete_category_rows"] += 1
        matches = []
        for row in rows:
            if str(row["major_category"]) != candidate.major_category:
                continue
            rule_categories = tuple(
                str(row[f"category_level_{index}"]) for index in range(1, 5)
            )
            score = score_service_fee_category_rule(rule_categories, category)
            if score is not None:
                matches.append(
                    (score, str(row["effective_from"]), int(row["id"]), row)
                )
        matched_item = max(matches, default=None, key=lambda item: item[:3])
        matched = None if matched_item is None else matched_item[3]
        if matched is None:
            counts["service_fee_unconfigured_rows"] += 1
        elif bool(matched["enabled"]):
            counts["service_fee_matched_rows"] += 1
        else:
            counts["service_fee_disabled_rows"] += 1
    return counts


def _enrich_major_categories(
    database_path: Path, candidates: list[MappingCandidate]
) -> list[MappingCandidate]:
    with connect_database(database_path) as connection:
        rows = connection.execute(
            """
            SELECT DISTINCT category_level_1, major_category
            FROM basic_service_fee_rates
            WHERE archived = 0
              AND category_level_1 <> '' AND category_level_1 <> '/'
              AND major_category <> '' AND major_category <> '/'
            """
        ).fetchall()
    grouped: dict[str, set[str]] = {}
    for row in rows:
        grouped.setdefault(str(row["category_level_1"]), set()).add(
            str(row["major_category"])
        )
    enriched: list[MappingCandidate] = []
    for candidate in candidates:
        options = grouped.get(candidate.category_level_1, set())
        if len(options) == 1:
            enriched.append(
                replace(
                    candidate,
                    major_category=next(iter(options)),
                    major_category_unconfigured_reason="",
                )
            )
        else:
            enriched.append(
                replace(
                    candidate,
                    major_category="",
                    major_category_unconfigured_reason=(
                        "经营大类无法唯一确定" if len(options) > 1 else "经营大类未配置"
                    ),
                )
            )
    return enriched


@management_locked
def apply_mapping_import(
    database_path: str | Path,
    preview: MappingImportPreview,
    *,
    imported_at: str,
) -> MappingImportResult:
    if not preview.ready_to_import:
        raise ValueError("预览存在重复或无效标识，禁止导入")

    verified = tuple(
        candidate
        for candidate in preview.candidates
        if candidate.status == "verified" and candidate.jushuitan_sku_id is not None
    )
    database = Path(database_path).resolve()
    with connect_database(database) as connection:
        connection.commit()
        enabled_sku_ids = {
            str(row[0])
            for row in connection.execute(
                "SELECT sku_id FROM products WHERE enabled = 1"
            )
        }
        unavailable = sorted(
            {
                candidate.jushuitan_sku_id
                for candidate in verified
                if candidate.jushuitan_sku_id not in enabled_sku_ids
            }
        )
        if unavailable:
            raise ValueError(
                f"导入前复核发现 {len(unavailable)} 个聚水潭商品已不可用，整批已回滚"
            )

        existing = {
            str(row["douyin_sku_id"]): str(row["jushuitan_sku_id"])
            for row in connection.execute(
                "SELECT douyin_sku_id, jushuitan_sku_id FROM douyin_sku_mappings"
            )
        }
        conflicts = [
            candidate.douyin_sku_id
            for candidate in verified
            if candidate.douyin_sku_id in existing
            and existing[candidate.douyin_sku_id] != candidate.jushuitan_sku_id
        ]
        if conflicts:
            raise ValueError(
                f"发现 {len(conflicts)} 个抖店 SKUID 映射冲突，整批已回滚"
            )

        new_candidates = [
            candidate for candidate in verified if candidate.douyin_sku_id not in existing
        ]
        existing_rows = len(verified) - len(new_candidates)
        if not new_candidates:
            previous = connection.execute(
                """
                SELECT id, inserted_rows, existing_rows
                FROM mapping_import_batches
                WHERE source_sha256 = ? AND status = 'applied'
                ORDER BY id DESC
                LIMIT 1
                """,
                (preview.source_sha256,),
            ).fetchone()
            if previous is not None:
                return MappingImportResult(
                    batch_id=int(previous["id"]),
                    inserted_rows=int(previous["inserted_rows"]),
                    existing_rows=int(previous["existing_rows"]),
                    reused_batch=True,
                )

        cursor = connection.execute(
            """
            INSERT INTO mapping_import_batches (
                source_file, source_sha256, total_rows, verified_rows, pending_rows,
                inserted_rows, existing_rows, status, imported_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'applied', ?)
            """,
            (
                preview.source_file,
                preview.source_sha256,
                preview.total_rows,
                preview.verified_rows,
                preview.pending_rows,
                len(new_candidates),
                existing_rows,
                imported_at,
            ),
        )
        batch_id = int(cursor.lastrowid)
        connection.executemany(
            """
            INSERT INTO douyin_sku_mappings (
                douyin_sku_id, douyin_product_id, merchant_sku_code,
                jushuitan_sku_id, product_name, specification, weight_kg,
                weight_source, weight_unconfigured_reason, category_level_1,
                category_level_2, category_level_3, category_level_4,
                major_category, major_category_unconfigured_reason, status,
                import_batch_id, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'verified', ?, ?, ?)
            """,
            (
                (
                    candidate.douyin_sku_id,
                    candidate.douyin_product_id,
                    candidate.merchant_sku_code,
                    candidate.jushuitan_sku_id,
                    candidate.product_name,
                    candidate.specification,
                    candidate.weight_kg,
                    candidate.weight_source,
                    candidate.weight_unconfigured_reason,
                    candidate.category_level_1,
                    candidate.category_level_2,
                    candidate.category_level_3,
                    candidate.category_level_4,
                    candidate.major_category,
                    candidate.major_category_unconfigured_reason,
                    batch_id,
                    imported_at,
                    imported_at,
                )
                for candidate in new_candidates
            ),
        )
        return MappingImportResult(
            batch_id=batch_id,
            inserted_rows=len(new_candidates),
            existing_rows=existing_rows,
            reused_batch=False,
        )


def preview_shop_mapping_changes(
    database_path: str | Path,
    preview: MappingImportPreview,
    shop_name: str,
) -> ShopMappingChanges:
    shop = normalize_shop_name(shop_name)
    database = Path(database_path).resolve()
    uri = f"{database.as_uri()}?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as connection:
        connection.row_factory = sqlite3.Row
        rows = connection.execute(
            "SELECT douyin_sku_id, douyin_product_id, merchant_sku_code, "
            "jushuitan_sku_id, product_name, specification, weight_kg, weight_source, "
            "weight_unconfigured_reason, category_level_1, "
            "category_level_2, category_level_3, category_level_4, "
            "major_category, major_category_unconfigured_reason, shop_name "
            "FROM douyin_sku_mappings"
        ).fetchall()
    return _calculate_shop_changes(rows, preview, shop)


@management_locked
def apply_shop_mapping_import(
    database_path: str | Path,
    preview: MappingImportPreview,
    *,
    shop_name: str,
    imported_at: str,
) -> ShopMappingImportResult:
    if not preview.ready_to_import:
        raise ValueError("预览存在重复或无效标识，禁止导入")
    shop = normalize_shop_name(shop_name)
    verified = tuple(
        candidate
        for candidate in preview.candidates
        if candidate.status == "verified" and candidate.jushuitan_sku_id is not None
    )
    if not verified:
        raise ValueError("当前店铺没有可导入的逐字符匹配记录")

    database = Path(database_path).resolve()
    with connect_database(database) as connection:
        connection.execute("BEGIN IMMEDIATE")
        enabled_sku_ids = {
            str(row[0])
            for row in connection.execute(
                "SELECT sku_id FROM products WHERE enabled = 1"
            )
        }
        unavailable = sorted(
            candidate.jushuitan_sku_id
            for candidate in verified
            if candidate.jushuitan_sku_id not in enabled_sku_ids
        )
        if unavailable:
            raise ValueError(
                f"导入前复核发现 {len(unavailable)} 个聚水潭商品已不可用，整批已回滚"
            )
        current_rows = connection.execute(
            "SELECT douyin_sku_id, douyin_product_id, merchant_sku_code, "
            "jushuitan_sku_id, product_name, specification, weight_kg, weight_source, "
            "weight_unconfigured_reason, category_level_1, "
            "category_level_2, category_level_3, category_level_4, "
            "major_category, major_category_unconfigured_reason, shop_name "
            "FROM douyin_sku_mappings"
        ).fetchall()
        changes = _calculate_shop_changes(current_rows, preview, shop)
        if changes.conflicting_sku_ids:
            raise ValueError(
                f"发现 {len(changes.conflicting_sku_ids)} 个 SKUID 已属于其它店铺，整批已回滚"
            )
        if changes.current_shop_rows > 0 and changes.shared_shop_rows == 0:
            raise ValueError("当前文件与已选店铺没有共同 SKUID，疑似选错店铺，禁止替换")
        if not any((changes.new_rows, changes.changed_rows, changes.removed_rows, changes.legacy_rows)):
            raise ValueError("当前店铺商品库已是最新，无需重复导入")

        stamp = datetime.now().astimezone().strftime("%Y%m%d-%H%M%S-%f")
        backup_path = database.with_name(
            f"{database.stem}.before-shop-mapping-{stamp}{database.suffix}.bak"
        )
        backup_sha256 = backup_database(database, backup_path)

        cursor = connection.execute(
            """
            INSERT INTO mapping_import_batches (
                source_file, source_sha256, total_rows, verified_rows, pending_rows,
                inserted_rows, existing_rows, updated_rows, removed_rows, legacy_rows,
                shop_name, import_mode, status, imported_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'replace', 'applied', ?)
            """,
            (
                preview.source_file,
                preview.source_sha256,
                preview.total_rows,
                preview.verified_rows,
                preview.pending_rows,
                changes.new_rows,
                changes.existing_rows,
                changes.changed_rows,
                changes.removed_rows,
                changes.legacy_rows,
                shop,
                imported_at,
            ),
        )
        batch_id = int(cursor.lastrowid)
        incoming_ids = {candidate.douyin_sku_id for candidate in verified}
        removed_ids = sorted(
            str(row["douyin_sku_id"])
            for row in current_rows
            if str(row["shop_name"]) == shop
            and str(row["douyin_sku_id"]) not in incoming_ids
        )
        connection.executemany(
            "DELETE FROM douyin_sku_mappings "
            "WHERE shop_name = ? AND douyin_sku_id = ?",
            ((shop, sku_id) for sku_id in removed_ids),
        )
        for candidate in verified:
            connection.execute(
                """
                INSERT INTO douyin_sku_mappings (
                    douyin_sku_id, douyin_product_id, merchant_sku_code,
                    jushuitan_sku_id, product_name, specification, weight_kg,
                    weight_source, weight_unconfigured_reason, category_level_1,
                    category_level_2, category_level_3, category_level_4,
                    major_category, major_category_unconfigured_reason, status,
                    import_batch_id, created_at, updated_at, shop_name
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'verified', ?, ?, ?, ?)
                ON CONFLICT(douyin_sku_id) DO UPDATE SET
                    douyin_product_id = excluded.douyin_product_id,
                    merchant_sku_code = excluded.merchant_sku_code,
                    jushuitan_sku_id = excluded.jushuitan_sku_id,
                    product_name = excluded.product_name,
                    specification = excluded.specification,
                    weight_kg = excluded.weight_kg,
                    weight_source = excluded.weight_source,
                    weight_unconfigured_reason = excluded.weight_unconfigured_reason,
                    category_level_1 = excluded.category_level_1,
                    category_level_2 = excluded.category_level_2,
                    category_level_3 = excluded.category_level_3,
                    category_level_4 = excluded.category_level_4,
                    major_category = excluded.major_category,
                    major_category_unconfigured_reason = excluded.major_category_unconfigured_reason,
                    status = 'verified',
                    import_batch_id = excluded.import_batch_id,
                    updated_at = excluded.updated_at,
                    shop_name = excluded.shop_name
                """,
                (
                    candidate.douyin_sku_id,
                    candidate.douyin_product_id,
                    candidate.merchant_sku_code,
                    candidate.jushuitan_sku_id,
                    candidate.product_name,
                    candidate.specification,
                    candidate.weight_kg,
                    candidate.weight_source,
                    candidate.weight_unconfigured_reason,
                    candidate.category_level_1,
                    candidate.category_level_2,
                    candidate.category_level_3,
                    candidate.category_level_4,
                    candidate.major_category,
                    candidate.major_category_unconfigured_reason,
                    batch_id,
                    imported_at,
                    imported_at,
                    shop,
                ),
            )
        mark_review_candidates_resolved_by_import(
            connection,
            preview.candidates,
            shop_name=shop,
            resolved_at=imported_at,
            batch_id=batch_id,
        )
        upsert_mapping_review_candidates(
            connection,
            preview.candidates,
            shop_name=shop,
            source_file=preview.source_file,
            source_sha256=preview.source_sha256,
            captured_at=imported_at,
            import_batch_id=batch_id,
        )
        record_system_update(
            connection,
            update_key=f"mapping:{batch_id}",
            kind="mapping",
            reference_id=str(batch_id),
            completed_at=imported_at,
            backup_path=backup_path,
            backup_sha256=backup_sha256,
            details={
                "batch_id": batch_id,
                "shop_name": shop,
                "source_file": preview.source_file,
                "mode": "replace",
                "total_rows": preview.total_rows,
                "verified_rows": preview.verified_rows,
                "pending_rows": preview.pending_rows,
                "inserted_rows": changes.new_rows,
                "existing_rows": changes.existing_rows,
                "updated_rows": changes.changed_rows,
                "removed_rows": changes.removed_rows,
                "legacy_rows": changes.legacy_rows,
            },
        )
        return ShopMappingImportResult(
            batch_id=batch_id,
            inserted_rows=changes.new_rows,
            existing_rows=changes.existing_rows,
            updated_rows=changes.changed_rows,
            removed_rows=changes.removed_rows,
            legacy_rows=changes.legacy_rows,
            backup_path=backup_path,
        )


def normalize_shop_name(value: object) -> str:
    if not isinstance(value, str):
        raise ValueError("请选择或填写店铺名称")
    shop = value.strip()
    if (
        not shop
        or len(shop) > 80
        or shop == "历史导入"
        or any(ord(character) < 32 or ord(character) == 127 for character in shop)
    ):
        raise ValueError("店铺名称无效；请使用管理台中固定且唯一的店铺名称")
    return shop


def _calculate_shop_changes(
    rows, preview: MappingImportPreview, shop_name: str
) -> ShopMappingChanges:
    current = {str(row["douyin_sku_id"]): row for row in rows}
    verified = {
        candidate.douyin_sku_id: candidate
        for candidate in preview.candidates
        if candidate.status == "verified" and candidate.jushuitan_sku_id is not None
    }
    current_shop_ids = {
        sku_id for sku_id, row in current.items() if str(row["shop_name"]) == shop_name
    }
    new_rows = existing_rows = changed_rows = legacy_rows = 0
    conflicts: list[str] = []
    for sku_id, candidate in verified.items():
        row = current.get(sku_id)
        if row is None:
            new_rows += 1
            continue
        owner = str(row["shop_name"])
        if owner == "历史导入":
            legacy_rows += 1
            continue
        if owner != shop_name:
            conflicts.append(sku_id)
            continue
        current_values = (
            str(row["douyin_product_id"]),
            str(row["merchant_sku_code"]),
            str(row["jushuitan_sku_id"]),
            str(row["product_name"]),
            str(row["specification"]),
            None if row["weight_kg"] is None else str(row["weight_kg"]),
            str(row["weight_source"]),
            str(row["weight_unconfigured_reason"]),
            str(row["category_level_1"]),
            str(row["category_level_2"]),
            str(row["category_level_3"]),
            str(row["category_level_4"]),
            str(row["major_category"]),
            str(row["major_category_unconfigured_reason"]),
        )
        incoming_values = (
            candidate.douyin_product_id,
            candidate.merchant_sku_code,
            candidate.jushuitan_sku_id,
            candidate.product_name,
            candidate.specification,
            candidate.weight_kg,
            candidate.weight_source,
            candidate.weight_unconfigured_reason,
            candidate.category_level_1,
            candidate.category_level_2,
            candidate.category_level_3,
            candidate.category_level_4,
            candidate.major_category,
            candidate.major_category_unconfigured_reason,
        )
        if current_values == incoming_values:
            existing_rows += 1
        else:
            changed_rows += 1
    return ShopMappingChanges(
        new_rows=new_rows,
        existing_rows=existing_rows,
        changed_rows=changed_rows,
        removed_rows=len(current_shop_ids - verified.keys()),
        legacy_rows=legacy_rows,
        current_shop_rows=len(current_shop_ids),
        shared_shop_rows=len(current_shop_ids & verified.keys()),
        conflicting_sku_ids=tuple(sorted(conflicts)),
    )


def rollback_mapping_import(
    database_path: str | Path,
    batch_id: int,
    *,
    rolled_back_at: str,
) -> int:
    database = Path(database_path).resolve()
    with connect_database(database) as connection:
        connection.execute("BEGIN IMMEDIATE")
        latest = connection.execute(
            "SELECT id, import_mode FROM mapping_import_batches WHERE status = 'applied' "
            "ORDER BY id DESC LIMIT 1"
        ).fetchone()
        if latest is None:
            raise ValueError("没有可回滚的映射导入批次")
        if int(latest["id"]) != batch_id:
            raise ValueError("只允许回滚最近一个仍生效的映射导入批次")
        if str(latest["import_mode"]) == "replace":
            raise ValueError("店铺整表替换请使用该批次生成的数据库备份恢复")
        deleted = connection.execute(
            "DELETE FROM douyin_sku_mappings WHERE import_batch_id = ?",
            (batch_id,),
        ).rowcount
        connection.execute(
            """
            UPDATE mapping_import_batches
            SET status = 'rolled_back', rolled_back_at = ?
            WHERE id = ?
            """,
            (rolled_back_at, batch_id),
        )
        return int(deleted)


def _load_enabled_sku_ids(database: Path) -> set[str]:
    uri = f"{database.as_uri()}?mode=ro"
    with closing(sqlite3.connect(uri, uri=True)) as connection:
        try:
            return {
                str(row[0])
                for row in connection.execute(
                    "SELECT sku_id FROM products WHERE enabled = 1"
                )
            }
        except sqlite3.Error as exc:
            raise ValueError("SQLite 缺少可读取的 products 商品表") from exc


def _identifier(value: object, field: str, source_row: int) -> str:
    if not isinstance(value, str):
        raise ValueError(f"第 {source_row} 行 {field} 必须按文本保存")
    if (
        not value.strip()
        or any(ord(character) < 32 or ord(character) == 127 for character in value)
    ):
        raise ValueError(f"第 {source_row} 行 {field} 无效")
    return value


def _header_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _display_text(value: object) -> str:
    return "" if value is None else str(value)


def _category_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="预览、导入或回滚抖店 SKU 批量映射")
    parser.add_argument("--database", required=True, help="已有 products 表的 SQLite 路径")
    parser.add_argument("--source", help="抖店商品导出 XLSX 路径")
    parser.add_argument("--confirmed-sha256", help="确认导入时必须逐字符提供预览哈希")
    parser.add_argument("--batch-id", type=int, help="回滚时指定最近一个生效批次 ID")
    parser.add_argument("command", choices=("preview", "import", "rollback"))
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    now = datetime.now().astimezone().replace(microsecond=0).isoformat()
    if args.command == "rollback":
        if args.batch_id is None:
            print("回滚失败：必须提供 --batch-id")
            return 1
        try:
            deleted = rollback_mapping_import(
                args.database, args.batch_id, rolled_back_at=now
            )
        except (OSError, ValueError) as exc:
            print(f"回滚失败：{exc}")
            return 1
        print(f"映射批次 {args.batch_id} 已回滚：删除新增映射 {deleted} 条")
        return 0

    if not args.source:
        print("预览失败：必须提供 --source")
        return 1
    try:
        preview = preview_mapping_import(args.database, args.source)
    except (OSError, ValueError) as exc:
        print(f"预览失败：{exc}")
        return 1
    print(
        "映射预览完成："
        f"总行={preview.total_rows}，唯一SKUID={preview.unique_douyin_skus}，"
        f"逐字符命中={preview.verified_rows}，待核验={preview.pending_rows}，"
        f"重复={preview.duplicate_rows}，无效={preview.invalid_rows}"
    )
    print(f"来源文件：{preview.source_file}")
    print(f"来源 SHA-256：{preview.source_sha256}")
    print("状态：可进入确认导入" if preview.ready_to_import else "状态：存在硬错误，禁止导入")
    if args.command == "preview":
        return 0 if preview.ready_to_import else 2
    if not preview.ready_to_import:
        return 2
    if args.confirmed_sha256 != preview.source_sha256:
        print("导入失败：--confirmed-sha256 与本次预览哈希不一致")
        return 1
    try:
        result = apply_mapping_import(args.database, preview, imported_at=now)
    except (OSError, ValueError) as exc:
        print(f"导入失败：{exc}")
        return 1
    reused = "（同一来源已导入，本次未重复写入）" if result.reused_batch else ""
    print(
        f"映射批次 {result.batch_id} 导入完成：新增={result.inserted_rows}，"
        f"已存在={result.existing_rows}{reused}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
