from __future__ import annotations

import argparse
import hashlib
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from .database import (
    connect_database,
    copy_mapping_state,
    count_products,
    set_sync_state,
    upsert_products,
)
from .excel import export_enabled_products, validate_excel
from .models import Product, normalize_decimal
from .locks import management_locked
from .sync import _create_temporary_file, _publish_pair
from .update_rollback import backup_database, record_system_update
from .xlsx_safety import MAX_IMPORT_ROWS, validate_xlsx_container


SHANGHAI = ZoneInfo("Asia/Shanghai")
ORDINARY_HEADERS = (
    "商品编码",
    "商品名称",
    "成本价",
    "商品状态",
    "创建时间",
    "修改时间",
)
COMBINATION_HEADERS = (
    "组合商品编码",
    "组合商品名称",
    "组合成本价",
    "商品状态",
    "创建时间",
    "修改时间",
    "商品编码",
    "数量",
    "子商品成本价",
)
EXCEL_ESCAPE_PATTERN = re.compile(r"_x([0-9A-Fa-f]{4})_")


@dataclass(frozen=True)
class CatalogImportPreview:
    ordinary_path: Path
    combination_path: Path
    ordinary_sha256: str
    combination_sha256: str
    ordinary_rows: int
    combination_rows: int
    ordinary_products: int
    combination_products: int
    merged_products: int
    empty_cost_products: int
    resolved_overlaps: tuple[str, ...]
    conflicting_overlaps: tuple[str, ...]
    combination_cost_mismatches: tuple[str, ...]
    products: tuple[Product, ...]

    @property
    def ready_to_import(self) -> bool:
        return not self.conflicting_overlaps


@dataclass(frozen=True)
class CatalogImportResult:
    total_count: int
    enabled_count: int
    empty_cost_products: int
    backup_path: Path | None
    completed_at: str


def preview_catalog_import(
    ordinary_path: str | Path, combination_path: str | Path
) -> CatalogImportPreview:
    ordinary = Path(ordinary_path).resolve()
    combination = Path(combination_path).resolve()
    if not ordinary.is_file():
        raise FileNotFoundError(f"普通商品表不存在：{ordinary}")
    if not combination.is_file():
        raise FileNotFoundError(f"组合商品表不存在：{combination}")
    validate_xlsx_container(ordinary, "普通商品表")
    validate_xlsx_container(combination, "组合商品表")

    ordinary_rows, ordinary_products = _load_ordinary_products(ordinary)
    (
        combination_rows,
        combination_products,
        combination_cost_mismatches,
    ) = _load_combination_products(combination)
    if not ordinary_products:
        raise ValueError("普通商品表没有商品数据，禁止更新")
    if not combination_products:
        raise ValueError("组合商品表没有商品数据，禁止更新")
    ordinary_date = _workbook_created_date(ordinary)
    combination_date = _workbook_created_date(combination)
    if (
        ordinary_date is not None
        and combination_date is not None
        and ordinary_date != combination_date
    ):
        raise ValueError(
            "普通商品表与组合商品表必须来自同一个导出日期："
            f"普通表 {ordinary_date.isoformat()}，组合表 {combination_date.isoformat()}"
        )
    merged = dict(ordinary_products)
    resolved: list[str] = []
    conflicts: list[str] = []

    for sku_id, combination_product in combination_products.items():
        ordinary_product = merged.get(sku_id)
        if ordinary_product is None:
            merged[sku_id] = combination_product
            continue
        if _can_combination_fill_missing_cost(ordinary_product, combination_product):
            merged[sku_id] = combination_product
            resolved.append(sku_id)
            continue
        conflicts.append(sku_id)

    products = tuple(merged[sku_id] for sku_id in sorted(merged))
    return CatalogImportPreview(
        ordinary_path=ordinary,
        combination_path=combination,
        ordinary_sha256=_sha256(ordinary),
        combination_sha256=_sha256(combination),
        ordinary_rows=ordinary_rows,
        combination_rows=combination_rows,
        ordinary_products=len(ordinary_products),
        combination_products=len(combination_products),
        merged_products=len(products),
        empty_cost_products=sum(product.cost_price is None for product in products),
        resolved_overlaps=tuple(sorted(resolved)),
        conflicting_overlaps=tuple(sorted(conflicts)),
        combination_cost_mismatches=tuple(sorted(combination_cost_mismatches)),
        products=products,
    )


@management_locked
def apply_catalog_import(
    database_path: str | Path,
    output_excel_path: str | Path,
    preview: CatalogImportPreview,
    *,
    confirmed_ordinary_sha256: str,
    confirmed_combination_sha256: str,
    imported_at: str,
) -> CatalogImportResult:
    if (
        preview.ordinary_products <= 0
        or preview.combination_products <= 0
        or preview.merged_products <= 0
    ):
        raise ValueError("两份商品表必须都包含商品数据，禁止导入空快照")
    if not preview.ready_to_import:
        raise ValueError("两份商品表存在未解决的重复编码冲突，禁止导入")
    if (
        confirmed_ordinary_sha256 != preview.ordinary_sha256
        or confirmed_combination_sha256 != preview.combination_sha256
        or _sha256(preview.ordinary_path) != preview.ordinary_sha256
        or _sha256(preview.combination_path) != preview.combination_sha256
    ):
        raise ValueError("两份商品表的 SHA-256 与预览不一致，禁止导入")

    completed_at = datetime.fromisoformat(imported_at).isoformat()
    database = Path(database_path).resolve()
    output_excel = Path(output_excel_path).resolve()
    database.parent.mkdir(parents=True, exist_ok=True)
    output_excel.parent.mkdir(parents=True, exist_ok=True)
    staged_database = _create_temporary_file(
        database.parent,
        prefix=f".{database.name}.",
        suffix=f".staged{database.suffix}",
    )
    staged_excel = _create_temporary_file(
        output_excel.parent,
        prefix=f".{output_excel.name}.",
        suffix=f".staged{output_excel.suffix}",
    )
    backup_path: Path | None = None
    backup_sha256: str | None = None
    try:
        if database.is_file():
            stamp = datetime.fromisoformat(completed_at).strftime("%Y%m%d-%H%M%S")
            backup_path = database.with_name(
                f"{database.stem}.before-paired-excel-import-{stamp}{database.suffix}.bak"
            )
            backup_sha256 = backup_database(database, backup_path)
        with connect_database(staged_database) as connection:
            if database.is_file():
                copy_mapping_state(database, connection)
            upsert_products(connection, preview.products, synced_at=completed_at)
            total_count = count_products(connection)
            enabled_count = count_products(connection, enabled_only=True)
            if total_count != preview.merged_products:
                raise ValueError("写入 SQLite 的商品数与预览不一致")
            set_sync_state(
                connection,
                {
                    "last_successful_sync": completed_at,
                    "last_mode": "excel_pair_import",
                    "last_fetched_count": str(total_count),
                    "last_enabled_count": str(enabled_count),
                    "last_ordinary_sha256": preview.ordinary_sha256,
                    "last_combination_sha256": preview.combination_sha256,
                    "last_ordinary_count": str(preview.ordinary_products),
                    "last_combination_count": str(preview.combination_products),
                },
            )
            if backup_path is not None and backup_sha256 is not None:
                record_system_update(
                    connection,
                    update_key=f"catalog:{completed_at}",
                    kind="catalog",
                    reference_id=completed_at,
                    completed_at=completed_at,
                    backup_path=backup_path,
                    backup_sha256=backup_sha256,
                    details={
                        "mode": "excel_pair_import",
                        "enabled_count": enabled_count,
                        "ordinary_count": preview.ordinary_products,
                        "combination_count": preview.combination_products,
                    },
                )
            connection.commit()
            exported_count = export_enabled_products(connection, staged_excel)
        if exported_count != enabled_count:
            raise ValueError("Excel 导出行数与 SQLite 启用商品数不一致")
        validate_excel(staged_excel, expected_rows=enabled_count)

        _publish_pair((staged_database, database), (staged_excel, output_excel))
    finally:
        staged_database.unlink(missing_ok=True)
        staged_excel.unlink(missing_ok=True)

    return CatalogImportResult(
        total_count=total_count,
        enabled_count=enabled_count,
        empty_cost_products=preview.empty_cost_products,
        backup_path=backup_path,
        completed_at=completed_at,
    )


def _load_ordinary_products(path: Path) -> tuple[int, dict[str, Product]]:
    rows = _read_rows(path, ORDINARY_HEADERS, "普通商品表")
    products: dict[str, Product] = {}
    row_count = 0
    for source_row, row in rows:
        row_count += 1
        sku_id = _identifier(row["商品编码"], "商品编码", source_row)
        if sku_id in products:
            raise ValueError(f"普通商品表第 {source_row} 行商品编码重复：{sku_id}")
        products[sku_id] = Product(
            sku_id=sku_id,
            autoid=None,
            name=_text(row["商品名称"]),
            cost_price=normalize_decimal(row["成本价"], sku_id=sku_id),
            enabled=_enabled(row["商品状态"], sku_id),
            created=_optional_text(row["创建时间"]),
            modified=_optional_text(row["修改时间"]),
        )
    return row_count, products


def _load_combination_products(
    path: Path,
) -> tuple[int, dict[str, Product], tuple[str, ...]]:
    rows = _read_rows(path, COMBINATION_HEADERS, "组合商品表")
    grouped: dict[str, list[tuple[int, dict[str, object]]]] = defaultdict(list)
    row_count = 0
    for source_row, row in rows:
        row_count += 1
        sku_id = _identifier(row["组合商品编码"], "组合商品编码", source_row)
        grouped[sku_id].append((source_row, row))

    products: dict[str, Product] = {}
    cost_mismatches: list[str] = []
    for sku_id, group in grouped.items():
        candidates = [
            Product(
                sku_id=sku_id,
                autoid=None,
                name=_text(row["组合商品名称"]),
                cost_price=normalize_decimal(row["组合成本价"], sku_id=sku_id),
                enabled=_enabled(row["商品状态"], sku_id),
                created=_optional_text(row["创建时间"]),
                modified=_optional_text(row["修改时间"]),
            )
            for _, row in group
        ]
        first = candidates[0]
        if any(candidate != first for candidate in candidates[1:]):
            source_rows = ", ".join(str(source_row) for source_row, _ in group)
            raise ValueError(f"组合商品 {sku_id} 的主数据不一致，来源行：{source_rows}")
        products[sku_id] = first
        if _combination_cost_mismatch(first, group):
            cost_mismatches.append(sku_id)
    return row_count, products, tuple(cost_mismatches)


def _combination_cost_mismatch(
    product: Product, group: list[tuple[int, dict[str, object]]]
) -> bool:
    if product.cost_price is None:
        return False
    total = Decimal("0")
    for _, row in group:
        try:
            quantity = Decimal(str(row["数量"]))
            child_cost = Decimal(str(row["子商品成本价"]))
        except (InvalidOperation, TypeError, ValueError):
            return True
        if not quantity.is_finite() or not child_cost.is_finite():
            return True
        if quantity < 0 or child_cost < 0:
            return True
        total += quantity * child_cost
    return total != Decimal(product.cost_price)


def _read_rows(path: Path, required_headers: tuple[str, ...], label: str):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError(f"{label}必须且只能包含一个工作表")
        worksheet = workbook[workbook.sheetnames[0]]
        worksheet.reset_dimensions()
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration as exc:
            raise ValueError(f"{label}为空") from exc
        headers = tuple(_text(value).strip() for value in raw_headers)
        duplicates = [
            header for header, count in Counter(headers).items() if header and count > 1
        ]
        if duplicates:
            raise ValueError(f"{label}存在重复列：{', '.join(duplicates)}")
        missing = [header for header in required_headers if header not in headers]
        if missing:
            raise ValueError(f"{label}缺少必填列：{', '.join(missing)}")
        indexes = {header: headers.index(header) for header in required_headers}
        materialized = []
        for source_row, values in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if len(materialized) >= MAX_IMPORT_ROWS:
                raise ValueError(f"{label}数据行超过 {MAX_IMPORT_ROWS:,} 行安全上限")
            materialized.append(
                (
                    source_row,
                    {header: values[index] for header, index in indexes.items()},
                )
            )
        return materialized
    finally:
        workbook.close()


def _can_combination_fill_missing_cost(
    ordinary: Product, combination: Product
) -> bool:
    return (
        ordinary.cost_price is None
        and combination.cost_price is not None
        and ordinary.enabled == combination.enabled
    )


def _identifier(value: object, field: str, source_row: int) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"第 {source_row} 行 {field} 不是有效文本编码")
    return EXCEL_ESCAPE_PATTERN.sub(
        lambda match: chr(int(match.group(1), 16)), value
    ).strip()


def _enabled(value: object, sku_id: str) -> int:
    normalized = _text(value).strip()
    if normalized == "启用":
        return 1
    if normalized in {"禁用", "停用"}:
        return -1
    raise ValueError(f"商品 {sku_id} 的商品状态无效：{normalized or '空'}")


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _optional_text(value: object) -> str | None:
    return None if value in (None, "") else str(value)


def _workbook_created_date(path: Path) -> date | None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        created = workbook.properties.created
        return None if created is None else created.date()
    finally:
        workbook.close()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="配对预览并导入聚水潭普通/组合商品表")
    parser.add_argument("--database", required=True, help="目标 SQLite 路径")
    parser.add_argument("--output-excel", required=True, help="启用商品成本 Excel 路径")
    parser.add_argument("--ordinary", required=True, help="普通商品导出表")
    parser.add_argument("--combination", required=True, help="组合商品导出表")
    parser.add_argument("--confirmed-ordinary-sha256")
    parser.add_argument("--confirmed-combination-sha256")
    parser.add_argument("command", choices=("preview", "import"))
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        preview = preview_catalog_import(args.ordinary, args.combination)
        print(
            f"预览：普通行={preview.ordinary_rows}，普通商品={preview.ordinary_products}，"
            f"组合明细={preview.combination_rows}，组合商品={preview.combination_products}，"
            f"合并商品={preview.merged_products}，空成本={preview.empty_cost_products}，"
            f"安全补全重复={len(preview.resolved_overlaps)}，"
            f"未解决冲突={len(preview.conflicting_overlaps)}"
        )
        print(f"普通表 SHA-256：{preview.ordinary_sha256}")
        print(f"组合表 SHA-256：{preview.combination_sha256}")
        if preview.conflicting_overlaps:
            print(f"冲突编码：{', '.join(preview.conflicting_overlaps)}")
        if args.command == "preview":
            return 0 if preview.ready_to_import else 1
        if not args.confirmed_ordinary_sha256 or not args.confirmed_combination_sha256:
            raise ValueError("导入必须提供两份预览文件的 SHA-256")
        result = apply_catalog_import(
            args.database,
            args.output_excel,
            preview,
            confirmed_ordinary_sha256=args.confirmed_ordinary_sha256,
            confirmed_combination_sha256=args.confirmed_combination_sha256,
            imported_at=datetime.now(SHANGHAI).replace(microsecond=0).isoformat(),
        )
        print(
            f"导入成功：总商品={result.total_count}，启用={result.enabled_count}，"
            f"空成本={result.empty_cost_products}，完成时间={result.completed_at}"
        )
        if result.backup_path:
            print(f"导入前备份：{result.backup_path}")
        return 0
    except (OSError, ValueError) as exc:
        print(f"执行失败：{exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
