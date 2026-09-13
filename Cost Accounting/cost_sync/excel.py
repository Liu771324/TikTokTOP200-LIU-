from __future__ import annotations

import sqlite3
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


SHEET_NAME = "启用商品成本价"
HEADERS = ("商品编码", "商品名称", "成本价", "最后同步时间")


def export_enabled_products(connection: sqlite3.Connection, path: str | Path) -> int:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = SHEET_NAME
    worksheet.append(HEADERS)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")

    rows = connection.execute(
        """
        SELECT sku_id, name, cost_price, synced_at
        FROM products
        WHERE enabled = 1
        ORDER BY sku_id
        """
    )
    count = 0
    for row in rows:
        synced_at = datetime.fromisoformat(row["synced_at"])
        price = None if row["cost_price"] is None else float(Decimal(row["cost_price"]))
        worksheet.append((row["sku_id"], row["name"], price, synced_at.replace(tzinfo=None)))
        count += 1
        row_number = count + 1
        worksheet.cell(row_number, 1).number_format = "@"
        worksheet.cell(row_number, 3).number_format = "0.00########"
        worksheet.cell(row_number, 4).number_format = "yyyy-mm-dd hh:mm:ss"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:D{max(1, worksheet.max_row)}"
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 50
    worksheet.column_dimensions["C"].width = 16
    worksheet.column_dimensions["D"].width = 22
    workbook.save(path)
    return count


def validate_excel(path: str | Path, *, expected_rows: int) -> None:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError("Excel 缺少目标工作表")
        worksheet = workbook[SHEET_NAME]
        headers = tuple(cell.value for cell in worksheet[1])
        if headers != HEADERS:
            raise ValueError("Excel 表头不符合规范")
        if worksheet.max_row - 1 != expected_rows:
            raise ValueError("Excel 数据行数与 SQLite 启用商品数不一致")
        for cell in worksheet["A"][1:]:
            if cell.number_format != "@":
                raise ValueError("Excel 商品编码列未按文本格式保存")
    finally:
        workbook.close()
