"""Offline protocol and parsing regressions for the Python business engine."""
import copy
import json
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from toolbox.engines.business import adapter, parser, transport


CATEGORY = {"displayPath": "一级 / 二级 / 三级", "segments": ["一级", "二级", "三级"], "valuePath": ["1", "2", "3"]}
SECOND = {"displayPath": "其他", "segments": ["其他"], "valuePath": ["4"]}


def row(cid, amount=900000, band="¥75万-¥100万", rate=0.2):
    return {"clue_detail": {"clue_id": cid, "name": "商品" + str(cid), "category_path": ["一级", "二级", "三级"]},
            "clue_indicator": {"pay_amount_ind": amount, "pay_amount_ind_range": band, "pay_amount_ind_30d_rate": rate},
            "cookie": "TOP_SECRET_MUST_NOT_BE_SAVED"}


def indicator(cid, hb=0.2, **extra):
    return parser.parse_indicator({"code": 0, "data": {"clue_id": cid, "pay_amount": 1200, "pay_amount_range": "¥1000-¥2500", "pay_amount_hb": hb, **extra},
                                   "cookie": "TOP_SECRET_MUST_NOT_BE_SAVED"}, str(cid))


class Cancelled(Exception):
    code = "CANCELLED"


class FakeApi:
    def __init__(self, context):
        self.context, self.events, self.registrations = context, [], []
        self.control = SimpleNamespace(cancelled=False)
        self.cancel_after_successes = None
        self.cancel_on_next_page = False
        self.pause, self.paused, self.resume = False, threading.Event(), threading.Event()
        self.after_wait = None
        self.pause_after_successes = None

    def emit(self, kind, data):
        self.events.append((kind, copy.deepcopy(data)))

    def checkpoint(self, state):
        self.state = copy.deepcopy(state)
        path = Path(self.context["statePath"])
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps({"formatVersion": 1, "toolId": "business-center", "taskId": self.context["taskId"],
                                         "attempt": self.context["attempt"], "state": state}, ensure_ascii=False, allow_nan=False), encoding="utf-8")
        temporary.replace(path)
        if self.cancel_after_successes is not None and len(state["results"]) >= self.cancel_after_successes:
            self.control.cancelled = True
        if self.cancel_on_next_page and state["combinations"][0]["nextPage"] == 2:
            self.control.cancelled = True
        if self.pause_after_successes is not None and len(state["results"]) >= self.pause_after_successes:
            self.pause = True
            self.pause_after_successes = None

    def boundary(self):
        if self.control.cancelled:
            raise Cancelled()
        if self.pause:
            self.paused.set()
            if not self.resume.wait(3):
                raise AssertionError("offline pause did not resume")
            self.pause = False
        if self.control.cancelled:
            raise Cancelled()

    def wait_for_user(self, message):
        self.events.append(("wait", {"message": message}))
        if self.after_wait is None:
            raise AssertionError("unexpected wait for login")
        self.after_wait()

    def register(self, candidate):
        path = Path(self.context["dataRoot"]) / candidate["relpath"]
        assert path.is_file() and path.stat().st_size > 0
        self.registrations.append(candidate)
        return {"artifactId": candidate["candidateId"], "registeredAttempt": self.context["attempt"]}


class FakeClient:
    def __init__(self, pages):
        self.pages, self.list_calls, self.indicator_calls = pages, [], []
        self.bad_ids, self.auth_once = set(), False
        self.on_indicator = None

    def fetch_list(self, category, number, size):
        self.list_calls.append((tuple(category["valuePath"]), number))
        if self.auth_once:
            self.auth_once = False
            raise parser.ImmediateStopError("请登录后继续")
        raw = self.pages.get((tuple(category["valuePath"]), number), self.pages.get(number, []))
        return parser.parse_list({"code": 0, "data": copy.deepcopy(raw)}, number)

    def fetch_indicator(self, cid):
        self.indicator_calls.append(cid)
        if self.on_indicator:
            self.on_indicator(cid)
        if cid in self.bad_ids:
            raise parser.CollectorError("指标 ID 错位")
        return indicator(cid)


class BusinessTests(unittest.TestCase):
    def test_fourth_level_platform_nodes_do_not_break_supported_category_scopes(self):
        payload = {"code": 0, "data": [{"value": 1, "label": "一级", "children": [
            {"value": 2, "label": "二级", "children": [{"value": 3, "label": "三级", "children": [
                {"value": 4, "label": "四级"}]}]}]}]}
        categories = parser.parse_categories(payload)
        self.assertEqual([item['valuePath'] for item in categories], [['1'], ['1', '2'], ['1', '2', '3']])

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="business-python-offline-", dir=Path(__file__).parent)
        root = Path(self.temporary.name)
        self.context = {"taskId": "offline-task", "attempt": 1, "dataRoot": str(root), "statePath": str(root / "state.json"),
                        "artifactStagingRoot": str(root / "staging"), "scratchRoot": str(root / "scratch"),
                        "browser": {"endpoint": "http://127.0.0.1:12345"}}
        self.api = FakeApi(self.context)
        self.parameters = {"categories": [CATEGORY], "minSales": 750000, "maxSales": None, "only30dGrowth": False}
        self.patches = [patch.object(adapter, "PAGE_INTERVAL", 0), patch.object(adapter, "RETRY_DELAYS", (0, 0))]
        for item in self.patches:
            item.start()

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.temporary.cleanup()

    def run_engine(self, client, api=None, context=None):
        session = SimpleNamespace(client=client, categories=lambda: [CATEGORY, SECOND], close=lambda: None)
        with patch.object(adapter, "ManagedSession", return_value=session):
            return adapter.run(self.parameters, context or self.context, api or self.api)

    def test_categories_official_three_level_and_invalid_ids(self):
        result = parser.parse_categories({"code": 0, "data": [{"value": 1, "label": "一级", "children": [{"value": 2, "label": "二级", "children": [{"value": 3, "label": "三级"}]}]}]})
        self.assertEqual(result[-1], CATEGORY)
        for payload in ({"data": [{"value": 0, "label": "bad"}]}, {"data": [{"value": 1, "label": "A"}, {"value": 1, "label": "A"}]}):
            with self.assertRaises(parser.CollectorError):
                parser.parse_categories(payload)

    def test_array_list_shape_empty_end_and_no_raw_auth(self):
        result = parser.parse_list({"code": 0, "data": [row(1)]}, 1)
        self.assertEqual(result["items"][0]["growth30dRate"], 0.2)
        self.assertNotIn("TOP_SECRET", json.dumps(result))
        self.assertEqual(parser.parse_list({"code": 0, "data": []}, 2)["items"], [])

    def test_amount_id_sort_and_band_validation(self):
        for rows in ([row(0)], [row(1), row(1)], [row(1, 800000), row(2, 900000)], [row(1, 5)], [row(1, float("nan"))]):
            with self.subTest(rows=rows), self.assertRaises(parser.CollectorError):
                parser.parse_list({"data": rows}, 1)
        self.assertEqual(parser.parse_amount_band("¥0-¥50")["min"], 0)
        self.assertEqual(parser.parse_amount_band("¥1亿+")["max"], float("inf"))

    def test_daily_six_classes_and_direction_validation(self):
        item = parser.parse_list({"data": [row(1)]}, 1)["items"][0]
        cases = [(indicator(1, -0.1), "下跌"), (indicator(1, 0), "持平"), (indicator(1), "上涨"),
                 (indicator(1, 1385.9357), "999.99%+"), (indicator(1, status_desc="增速快"), "增速快"),
                 (indicator(1, status_desc="暂无数据"), "无法确认")]
        for response, expected in cases:
            self.assertEqual(parser.classify_indicator(response, item)["classification"], expected)
        for response in (indicator(1, 0.2, pay_amount_hb_direction="down"), indicator(1, 0.2, status_desc="持平")):
            with self.assertRaises(parser.CollectorError):
                parser.classify_indicator(response, item)
        for payload in ({"data": {"clue_id": 999}}, {"data": {"clue_id": 0}}):
            with self.assertRaises(parser.CollectorError):
                parser.parse_indicator(payload, "1")
        self.assertEqual(indicator(1)["idBinding"], "response-id-verified")

    def test_real_indicator_schema_without_response_id_is_request_bound(self):
        response = parser.parse_indicator({"code": 0, "data": {"pay_amount": 1200, "pay_amount_range": "¥1000-¥2500", "pay_amount_hb": 0.2, "is_growth_rate_high": False}}, "1")
        self.assertEqual(response["idBinding"], "request-bound")
        self.assertNotIn("clue_id", response["data"])
        client = FakeClient({1: [row(1)]})
        with patch.object(client, "fetch_indicator", return_value=response):
            result = self.run_engine(client)
        self.assertEqual(result["state"], "completed")
        self.assertEqual(result["checks"]["requestBound"], 1)
        self.assertEqual(self.api.state["results"][0]["indicatorIdBinding"], "request-bound")

    def test_multi_category_bounds_report_and_read_results(self):
        self.parameters.update({"categories": [CATEGORY, SECOND], "maxSales": 1000000})
        client = FakeClient({1: [row(1, 1100000, "¥100万-¥250万"), row(2), row(3, 700000, "¥50万-¥75万")]})
        result = self.run_engine(client)
        self.assertEqual(result["state"], "completed")
        self.assertEqual(client.indicator_calls, ["2", "2"])
        self.assertTrue(result["coverageComplete"])
        self.assertEqual(len(self.api.registrations), 3)
        self.assertEqual({item["role"] for item in self.api.registrations}, {"report", "report-json", "report-markdown"})
        self.assertNotIn("TOP_SECRET", Path(self.context["statePath"]).read_text(encoding="utf-8"))
        preview = adapter.read_results({"offset": 1, "limit": 1}, self.context)
        self.assertEqual((preview["total"], len(preview["items"])), (2, 1))
        self.assertEqual(preview["items"][0]["productName"], "商品2")
        from openpyxl import load_workbook
        book = load_workbook(Path(self.context["artifactStagingRoot"]) / "business-report.xlsx", read_only=True)
        self.assertIn("覆盖与检查", book.sheetnames)
        self.assertEqual(book["全部数据"].max_row, 3)
        book.close()

    def test_only30d_growth_field_and_invalid_values(self):
        self.parameters["only30dGrowth"] = True
        client = FakeClient({1: [row(1, rate=-0.114), row(2, rate=0.1532), row(3, rate=0), row(4, rate=None), row(5, rate="999%"), row(6, rate=float("inf"))]})
        result = self.run_engine(client)
        self.assertEqual(client.indicator_calls, ["2"])
        self.assertEqual(result["state"], "completed_with_errors")
        self.assertFalse(result["coverageComplete"])
        self.assertEqual(result["stats"]["filteredBy30dGrowth"], 2)
        self.assertEqual(result["stats"]["invalid30dGrowth"], 3)
        self.assertEqual(len(adapter.read_results({"section": "unparseableItems"}, self.context)["items"]), 3)

    def test_stop_partial_resume_current_page_deduplicates(self):
        client = FakeClient({1: [row(number) for number in range(1, 7)]})
        self.api.cancel_after_successes = 3
        result = self.run_engine(client)
        self.assertEqual(result["state"], "cancelled")
        self.assertEqual(self.api.state["combinations"][0]["nextPage"], 1)
        self.assertEqual(len(self.api.state["results"]), 3)
        self.assertEqual(len(self.api.registrations), 3)
        self.context["attempt"] = 2
        resumed = FakeApi(self.context)
        self.assertEqual(self.run_engine(client, resumed)["state"], "completed")
        self.assertEqual(len(client.indicator_calls), 6)
        self.assertEqual(len(set(client.indicator_calls)), 6)
        self.assertEqual(len(resumed.state["results"]), 6)

    def test_cancelled_pending_batch_keeps_finished_results(self):
        client = FakeClient({1: [row(number) for number in range(1, 7)]})
        client.on_indicator = lambda cid: setattr(self.api.control, "cancelled", True)
        result = self.run_engine(client)
        self.assertEqual(result["state"], "cancelled")
        self.assertGreaterEqual(len(self.api.state["results"]), 1)
        self.assertEqual(len(self.api.state["results"]), len(client.indicator_calls))

    def test_resume_next_page_cross_page_sort_and_repetition(self):
        for bad_page, expected in (([row(1), row(2)], "重复"), ([row(3, 950000), row(4, 900000)], "降序"), ([row(2), row(3)], "跨页重复")):
            with self.subTest(expected=expected):
                Path(self.context["statePath"]).unlink(missing_ok=True)
                self.api = FakeApi(self.context)
                with patch.object(adapter, "PAGE_SIZE", 2):
                    result = self.run_engine(FakeClient({1: [row(1), row(2, 800000)], 2: bad_page}))
                self.assertEqual(result["state"], "completed_with_errors")
                self.assertIn(expected, self.api.state["failures"][-1]["reason"])
                self.assertEqual(len(self.api.state["results"]), 2)
                self.assertFalse(result["coverageComplete"])

    def test_resume_refuses_changed_pending_page(self):
        self.api.cancel_after_successes = 3
        client = FakeClient({1: [row(number) for number in range(1, 7)]})
        self.run_engine(client)
        client.pages[1][0] = row(99)
        self.context["attempt"] = 2
        resumed = FakeApi(self.context)
        result = self.run_engine(client, resumed)
        self.assertEqual(result["state"], "completed_with_errors")
        self.assertIn("数据已变化", resumed.state["failures"][-1]["reason"])
        self.assertEqual(len(resumed.state["results"]), 3)

    def test_retry_only_failed_indicator_preserves_success_categories(self):
        self.parameters["categories"] = [CATEGORY, SECOND]
        client = FakeClient({("1", "2", "3"): []})
        client.pages = {(tuple(CATEGORY["valuePath"]), 1): [row(1), row(2)], (tuple(SECOND["valuePath"]), 1): [row(3)]}
        client.bad_ids = {"2"}
        self.assertEqual(self.run_engine(client)["state"], "completed_with_errors")
        calls = list(client.list_calls)
        self.context.update({"attempt": 2, "retryFailed": True})
        client.bad_ids.clear()
        retried = FakeApi(self.context)
        self.assertEqual(self.run_engine(client, retried)["state"], "completed")
        self.assertEqual(client.list_calls, calls)
        self.assertEqual(client.indicator_calls.count("1"), 1)
        self.assertEqual(client.indicator_calls.count("3"), 1)
        self.assertEqual(len(retried.state["results"]), 3)

    def test_stopped_failure_retry_continues_remaining_ids_on_resume(self):
        client = FakeClient({1: [row(number) for number in range(1, 7)]})
        client.bad_ids = {str(number) for number in range(1, 7)}
        self.assertEqual(self.run_engine(client)['state'], 'completed_with_errors')
        self.context.update({'attempt': 2, 'retryFailed': True})
        client.bad_ids.clear()
        retry_api = FakeApi(self.context)
        retry_api.cancel_after_successes = 3
        self.assertEqual(self.run_engine(client, retry_api)['state'], 'cancelled')
        self.assertTrue(retry_api.state['retryingFailures'])
        self.assertEqual(len(retry_api.state['results']), 3)
        calls_before = len(client.indicator_calls)
        self.context.update({'attempt': 3})
        self.context.pop('retryFailed')
        resumed = FakeApi(self.context)
        self.assertEqual(self.run_engine(client, resumed)['state'], 'completed')
        self.assertEqual(len(client.indicator_calls) - calls_before, 3)
        self.assertEqual(len(resumed.state['results']), 6)
        self.assertEqual(len({item['clueId'] for item in resumed.state['results']}), 6)
        self.assertFalse(resumed.state['retryingFailures'])

    def test_retry_invalid_30d_field_only_original_id(self):
        self.parameters["only30dGrowth"] = True
        client = FakeClient({1: [row(1, rate=None), row(2)]})
        self.assertEqual(self.run_engine(client)["state"], "completed_with_errors")
        client.pages[1][0] = row(1, rate=0.4)
        self.context.update({"attempt": 2, "retryFailed": True})
        retried = FakeApi(self.context)
        result = self.run_engine(client, retried)
        self.assertEqual(result["state"], "completed")
        self.assertEqual(client.indicator_calls.count("2"), 1)
        self.assertEqual(client.indicator_calls.count("1"), 1)
        self.assertFalse(retried.state["unparseableItems"])

    def test_pause_blocks_new_requests_until_resume(self):
        self.api.pause = True
        client = FakeClient({1: [row(1)]})
        results, errors = [], []
        def run():
            try:
                results.append(self.run_engine(client))
            except BaseException as error:
                errors.append(error)
        thread = threading.Thread(target=run)
        thread.start()
        self.assertTrue(self.api.paused.wait(1))
        self.assertEqual(client.list_calls, [])
        self.assertEqual(client.indicator_calls, [])
        self.api.resume.set()
        thread.join(4)
        self.assertFalse(thread.is_alive())
        self.assertEqual(errors, [])
        self.assertEqual(results[0]["state"], "completed")

    def test_login_wait_preserves_progress_and_reconnects(self):
        client = FakeClient({1: [row(1)]})
        client.auth_once = True
        self.api.after_wait = lambda: None
        self.assertEqual(self.run_engine(client)["state"], "completed")
        self.assertEqual(sum(kind == "wait" for kind, data in self.api.events), 1)
        self.assertEqual(client.indicator_calls, ["1"])

    def test_legacy_read_only_and_parameter_identity(self):
        state = adapter.new_state(adapter.validate_parameters(self.parameters), self.context)
        state.pop("engine")
        state["results"] = [{"keyword": "旧商品", "salesAmount": "75万-100万", "categoryPath": "旧类目"}]
        self.api.checkpoint(state)
        self.assertTrue(adapter.read_results({}, self.context)["readOnly"])
        with self.assertRaisesRegex(parser.CollectorError, "旧 DOM"):
            self.run_engine(FakeClient({}))
        state = adapter.new_state(adapter.validate_parameters(self.parameters), self.context)
        self.api.checkpoint(state)
        self.parameters["minSales"] = 800000
        with self.assertRaisesRegex(parser.CollectorError, "不能更改"):
            self.run_engine(FakeClient({}))

    def test_transport_fixed_endpoints_headers_and_official_ids(self):
        client = transport.ApiClient("secret-cookie", "fake-ua")
        calls = []
        def fake_request(url, payload, headers=None):
            calls.append((url, payload, headers))
            return {"code": 0, "data": []}
        with patch.object(client, "request", side_effect=fake_request):
            client.fetch_list(CATEGORY, 2, 18)
        self.assertEqual(calls[0][1]["condition"]["categories"], [{"first_cid": 1, "second_cid": 2, "third_cid": 3}])
        self.assertEqual(calls[0][0], transport.LIST_URL)
        with self.assertRaisesRegex(parser.CollectorError, "不受支持"):
            client.request("https://attacker.invalid/", {})
        for endpoint in (None, "http://evil.invalid:9222", "http://127.0.0.1", "http://u:p@127.0.0.1:9222"):
            with self.assertRaises(parser.CollectorError):
                transport.ManagedSession({"browser": {"endpoint": endpoint}}, self.api)

    def test_short_page_requires_empty_confirmation_and_total_mismatch_is_partial(self):
        client = FakeClient({1: [row(1)], 2: [row(2)]})
        self.assertEqual(self.run_engine(client)["state"], "completed")
        self.assertEqual([number for ids, number in client.list_calls], [1, 2, 3])
        Path(self.context["statePath"]).unlink()
        self.api = FakeApi(self.context)
        client = FakeClient({1: {"list": [row(1)], "has_more": False, "total": 9}})
        result = self.run_engine(client)
        self.assertEqual(result["state"], "completed_with_errors")
        self.assertIn("总条数不一致", self.api.state["failures"][-1]["reason"])
        self.assertEqual(len(self.api.state["results"]), 1)

    def test_pause_after_batch_prevents_second_batch_requests(self):
        self.api.pause_after_successes = 3
        client = FakeClient({1: [row(number) for number in range(1, 7)]})
        results = []
        thread = threading.Thread(target=lambda: results.append(self.run_engine(client)))
        thread.start()
        self.assertTrue(self.api.paused.wait(2))
        self.assertEqual(len(client.indicator_calls), 3)
        self.api.resume.set()
        thread.join(4)
        self.assertFalse(thread.is_alive())
        self.assertEqual(results[0]["state"], "completed")
        self.assertEqual(len(client.indicator_calls), 6)

    def test_managed_cdp_reads_existing_page_only_and_filters_signed_url(self):
        operations = []
        signed = transport.BASE_URL + transport.CATEGORY_PATH + "?signature=TOP_SECRET"
        page = SimpleNamespace(url=transport.BASE_URL + transport.PAGE_PATH,
                               evaluate=lambda expression: "fake-ua" if expression == "navigator.userAgent" else ["https://evil.invalid" + transport.CATEGORY_PATH, signed])
        browser_context = SimpleNamespace(pages=[page], cookies=lambda origins: operations.append(("cookies", origins)) or [{"name": "session", "value": "TOP_SECRET"}])
        browser = SimpleNamespace(contexts=[browser_context])
        playwright = SimpleNamespace(chromium=SimpleNamespace(connect_over_cdp=lambda endpoint, timeout: operations.append(("connect", endpoint, timeout)) or browser),
                                     stop=lambda: operations.append(("disconnect",)))
        module = SimpleNamespace(sync_playwright=lambda: SimpleNamespace(start=lambda: playwright))
        with patch.dict(sys.modules, {"playwright.sync_api": module}):
            session = transport.ManagedSession(self.context, self.api)
            with patch.object(session.client, "request", return_value={"data": [{"value": 1, "label": "一级"}]}) as request:
                self.assertEqual(session.categories()[0]["valuePath"], ["1"])
                request.assert_called_once_with(signed, None)
            session.close()
        self.assertEqual(operations, [("connect", self.context["browser"]["endpoint"], 5000), ("cookies", [transport.BASE_URL]), ("disconnect",)])

    def test_transport_bounded_body_redirect_and_safe_errors(self):
        from unittest.mock import MagicMock
        session = MagicMock()
        session.__enter__.return_value = session
        response = MagicMock(status_code=200, headers={"content-type": "application/json"})
        response.__enter__.return_value = response
        response.raw.read1.side_effect = [b'{"code":0,"data":[]}', b""]
        session.request.return_value = response
        client = transport.ApiClient("TOP_SECRET", "fake-ua")
        with patch.object(transport.requests, "Session", return_value=session):
            self.assertEqual(client.request(transport.LIST_URL, {})["data"], [])
        kwargs = session.request.call_args.kwargs
        self.assertEqual(kwargs["timeout"], (2, 2))
        self.assertFalse(kwargs["allow_redirects"])
        self.assertFalse(session.trust_env)
        response.status_code = 302
        with patch.object(transport.requests, "Session", return_value=session), self.assertRaises(parser.ImmediateStopError):
            client.request(transport.LIST_URL, {})
        session.request.side_effect = RuntimeError("https://secret.invalid/?cookie=TOP_SECRET")
        with patch.object(transport.requests, "Session", return_value=session):
            with self.assertRaises(parser.CollectorError) as raised:
                client.request(transport.LIST_URL, {})
        self.assertNotIn("TOP_SECRET", str(raised.exception))

    def test_invalid_checkpoint_progress_is_not_overwritten(self):
        state = adapter.new_state(adapter.validate_parameters(self.parameters), self.context)
        state["combinations"][0]["nextPage"] = 12
        self.api.checkpoint(state)
        before = Path(self.context["statePath"]).read_bytes()
        with self.assertRaisesRegex(parser.CollectorError, "进度损坏"):
            self.run_engine(FakeClient({}))
        self.assertEqual(Path(self.context["statePath"]).read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
