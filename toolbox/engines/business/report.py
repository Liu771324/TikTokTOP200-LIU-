"""Write partial or complete business reports in the host staging directory."""
from __future__ import annotations

import json
import os
from pathlib import Path


def safe_cell(value):
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def markdown(state):
    stats = state["stats"]
    lines = ["# 商机中心采集报告", "", f"任务：{state['taskId']}", f"状态：{state['status']}",
             f"完整数据覆盖：{'已完成' if state['coverageComplete'] else '未完成（部分报告）'}",
             f"已扫描：{stats['scanned']}；金额达标：{stats['qualifying']}；成功：{stats['success']}；失败：{stats['failed']}",
             f"近30天增长筛掉：{stats['filteredBy30dGrowth']}；近30天增长无法解析：{stats['invalid30dGrowth']}",
             "", "金额筛选采用整个档位包含在最低/最高金额范围内的规则。",
             "近30天增长筛选采用官方列表 pay_amount_ind_30d_rate > 0；缺失或无效单独记录，不使用同比字段。",
             "结果的增长字段为今日相较昨日的环比；昨日成交档位由今日金额和环比反推，不是平台昨日原值。",
             "", "## 分类与检查", ""]
    lines.extend(f"- {name}：{count}" for name, count in stats["classifications"].items())
    response_ids = sum(item["indicatorIdBinding"] == "response-id-verified" for item in state["results"])
    request_ids = sum(item["indicatorIdBinding"] == "request-bound" for item in state["results"])
    lines.extend(["", "商品 ID、指标 ID 绑定、金额与档位、页内/跨页降序、重复页及分页结束条数均经过检查。",
                  f"指标返回 ID 并核对：{response_ids}；指标未返回 ID、按独立请求及 future 对原商品绑定：{request_ids}（未伪造响应 ID）。",
                  "", "## 类目覆盖", ""])
    for category in state["combinations"]:
        lines.append(f"- {category['category']['displayPath']}（{' / '.join(category['category']['valuePath'])}）："
                     f"{'列表范围已覆盖' if category['coverageComplete'] else '未完成'}；下一页 {category['nextPage']}；"
                     f"成功 {category['resultCount']}；失败 {category['failureCount']}；未处理达标项 "
                     f"{len(set(category['qualifyingIds']) - set(category['processedIds']))}")
        if category.get("error"):
            lines.append(f"  检查异常：{category['error']}")

    def clean(value):
        return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ").replace("\r", " ")

    lines.extend(["", "## 成功结果", "", "| 类目 | 商品 | 近30天成交 | 昨日（反推） | 今日 | 今日环比 | 分类 |",
                  "|---|---|---|---|---|---|---|"])
    for item in state["results"]:
        lines.append("| " + " | ".join(clean(item[key]) for key in ("categoryPath", "productName", "totalSales", "yesterday", "today", "growth", "classification")) + " |")
    lines.extend(["", "## 未解决失败", "", "| 类目 | 页码 | 商品 | 阶段 | 原因 | 尝试次数 |", "|---|---|---|---|---|---|"])
    for item in state["failures"] + state["unparseableItems"]:
        lines.append("| " + " | ".join(clean(item.get(key, "")) for key in ("categoryPath", "page", "productName", "stage", "reason", "attempts")) + " |")
    return "\n".join(lines) + "\n"


def excel(state, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "采集汇总"
    summary.append(["任务编号", "状态", "完整数据覆盖", "已扫描", "金额达标", "成功", "失败", "未完成类目"])
    stats = state["stats"]
    summary.append([state["taskId"], state["status"], "已完成" if state["coverageComplete"] else "未完成（部分报告）",
                    stats["scanned"], stats["qualifying"], stats["success"], stats["failed"], stats["totalCategories"] - stats["completedCategories"]])
    summary.append(["金额筛选", "完整金额档位包含在最低/最高范围内"])
    summary.append(["增长口径", "今日相较昨日环比，非近30天增长"])
    summary.append(["近30天筛选口径", "列表pay_amount_ind_30d_rate > 0", "筛掉", stats["filteredBy30dGrowth"], "无法解析", stats["invalid30dGrowth"]])
    summary.append(["昨日口径", "由今日精确金额及环比反推的档位，非平台原始值"])
    summary.append(["指标ID绑定", "响应ID核对", sum(item["indicatorIdBinding"] == "response-id-verified" for item in state["results"]),
                    "独立请求绑定（接口无返回ID）", sum(item["indicatorIdBinding"] == "request-bound" for item in state["results"])])
    summary.append(["分类", *stats["classifications"].keys()])
    summary.append(["数量", *stats["classifications"].values()])
    headers = ["类目", "商品ID", "商品", "近30天成交档位", "近30天精确金额", "昨日（反推）", "今日", "今日环比", "分类", "页码", "尝试次数"]
    for name, rows in (("全部数据", state["results"]), ("成交增长", [item for item in state["results"] if item["classification"] in ("上涨", "999.99%+", "增速快")])):
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for item in rows:
            sheet.append([safe_cell(item[key]) for key in ("categoryPath", "clueId", "productName", "totalSales", "pay30d", "yesterday", "today", "growth", "classification", "page", "attempts")])
    failures = workbook.create_sheet("失败记录")
    failures.append(["类目", "页码", "商品ID", "商品", "阶段", "原因", "尝试次数"])
    for item in state["failures"] + state["unparseableItems"]:
        failures.append([safe_cell(item.get(key, "")) for key in ("categoryPath", "page", "clueId", "productName", "stage", "reason", "attempts")])
    coverage = workbook.create_sheet("覆盖与检查")
    coverage.append(["类目", "官方ID路径", "覆盖", "下一页", "已检查页", "已扫描", "金额达标", "成功", "失败", "未处理达标", "检查异常"])
    for category in state["combinations"]:
        coverage.append([safe_cell(category["category"]["displayPath"]), " / ".join(category["category"]["valuePath"]),
                         "已覆盖" if category["coverageComplete"] else "未完成", category["nextPage"], len(category["pages"]), len(category["scannedIds"]),
                         len(category["qualifyingIds"]), category["resultCount"], category["failureCount"],
                         len(set(category["qualifyingIds"]) - set(category["processedIds"])), safe_cell(category.get("error") or "")])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17263B")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(58, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
    temporary = path.with_suffix(".tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def save_reports(state, context, api):
    root = Path(context["artifactStagingRoot"]).resolve()
    data_root = Path(context["dataRoot"]).resolve()
    root.relative_to(data_root)
    root.mkdir(parents=True, exist_ok=True)
    paths = {"excel": root / "business-report.xlsx", "json": root / "business-report.json", "markdown": root / "business-report.md"}
    for key, value in (("json", json.dumps(state, ensure_ascii=False, indent=2, allow_nan=False)), ("markdown", markdown(state))):
        path = paths[key]
        temporary = path.with_suffix(path.suffix + ".tmp")
        try:
            temporary.write_text(value, encoding="utf-8")
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
    excel(state, paths["excel"])
    registrations = {}
    for key, candidate_id, kind, role, display in (
        ("excel", "business-report", "excel", "report", "商机中心.xlsx"),
        ("json", "business-report-json", "text", "report-json", "商机中心.json"),
        ("markdown", "business-report-markdown", "text", "report-markdown", "商机中心.md"),
    ):
        registrations[key] = api.register({"candidateId": candidate_id, "kind": kind, "role": role,
                                           "relpath": paths[key].relative_to(data_root).as_posix(), "displayName": display,
                                           "replaceRole": True, "metadata": {"categories": len(state["combinations"]),
                                                                                "coverageComplete": state["coverageComplete"], "engine": state["engine"]}})
    return registrations
